# code : App_IA_Assistant_Gradio.py
"""Assistant Data AI : interface Gradio.

Lancement (depuis la racine du projet) :
    python -m app.App_IA_Assistant_Gradio
"""

from __future__ import annotations

import html
import logging
import re
import tempfile
import time
from decimal import Decimal
from functools import lru_cache
from pathlib import Path
from typing import Any

import gradio as gr
import pandas as pd
import plotly.graph_objects as go

from app.llm.openai_client import OpenAIClient
from app.pipeline.data_pipeline import DataPipeline

logger = logging.getLogger(__name__)

# ---------------------------------------------------------
# Constantes
# ---------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
IMAGE_DIR = BASE_DIR / "image"

LOGO_PATH = next(
    (
        path
        for path in (IMAGE_DIR / "logo-icon.png", IMAGE_DIR / "generated-image.png")
        if path.is_file()
    ),
    IMAGE_DIR / "logo-icon.png",
)
HAS_LOGO = LOGO_PATH.is_file()

EXAMPLES = [
    "Quel est le chiffre d'affaires par région ?",
    "Quels sont mes 5 produits les plus rentables ?",
    "Compare 2024 et 2023 : qu'est-ce qui a changé ?",
    "Analyse les ventes et identifie les tendances ou anomalies",
    "Prévois la marge pour les 6 prochains mois",
]

WELCOME = (
    "Bonjour 👋 Je suis votre assistant Data. "
    "Posez-moi simplement votre question en langage naturel "
    "et je vous aiderai à trouver l'information recherchée."
)

# Couleurs
INDIGO = "#4f46e5"
TEXT_DARK = "#1e1b4b"
ZEBRA = ("#ffffff", "#f5f3ff")
HEAT_LOW = (224, 231, 255)
HEAT_HIGH = (79, 70, 229)
MAX_STYLED_ROWS = 5000

TIME_WORDS = (
    "annee", "année", "year", "mois", "month", "jour", "day",
    "trimestre", "quarter", "semaine", "week", "date",
    "periode", "période",
)
NON_ADDITIVE_WORDS = ("taux", "moyenne", "ratio", "pourcent", "%", "prix", "pct", "avg")

LABEL_REPLACEMENTS = (
    (r"\bchiffre affaires\b", "chiffre d'affaires"),
    (r"\bquantite\b", "quantité"),
    (r"\bregion\b", "région"),
    (r"\bannee\b", "année"),
    (r"\bcategorie\b", "catégorie"),
    (r"\bcout\b", "coût"),
    (r"\bbenefice\b", "bénéfice"),
    (r"\bca\b", "CA"),
)

MONTHS_FR = (
    "janvier", "février", "mars", "avril", "mai", "juin",
    "juillet", "août", "septembre", "octobre", "novembre", "décembre",
)
MONTH_COLUMN_NAMES = ("mois", "month")

FOLLOW_UP_DIMENSIONS = (
    ("region", "région"),
    ("produit", "produit"),
    ("mois", "mois"),
    ("annee", "année"),
)

FORCE_LIGHT_JS = """
() => {
    const url = new URL(window.location);
    if (url.searchParams.get('__theme') !== 'light') {
        url.searchParams.set('__theme', 'light');
        window.location.href = url.href;
    }
}
"""
MAP_CLICK_JS = """
() => {
  if (window.__mapClickInstalled) return;
  window.__mapClickInstalled = true;
  const attach = () => {
    document.querySelectorAll('#chart .js-plotly-plot').forEach(gd => {
      if (gd.__clickBound || !gd.on) return;
      gd.__clickBound = true;
      gd.on('plotly_click', ev => {
        const p = ev.points && ev.points[0];
        if (!p) return;
        const name = p.location || (typeof p.x === 'string' ? p.x : p.y);
        const box = document.querySelector('#map-click textarea, #map-click input');
        if (!name || !box) return;
        box.value = String(name);
        box.dispatchEvent(new Event('input', { bubbles: true }));
      });
    });
  };
  new MutationObserver(attach).observe(document.body, { childList: true, subtree: true });
  attach();
}
"""
# Le CSS est dans app/style.css (à créer : voir les instructions).
CSS_PATH = BASE_DIR / "style.css"
CSS = CSS_PATH.read_text(encoding="utf-8") if CSS_PATH.is_file() else ""
if not CSS:
    logger.warning("style.css introuvable : l'interface s'affichera sans style personnalisé.")

THEME = gr.themes.Soft(
    primary_hue="indigo",
    neutral_hue="slate",
    font=[gr.themes.GoogleFont("Inter"), "system-ui", "sans-serif"],
).set(
    body_background_fill="#ffffff",
    body_background_fill_dark="#ffffff",
    background_fill_primary="#ffffff",
    background_fill_primary_dark="#ffffff",
    background_fill_secondary="#f8fafc",
    background_fill_secondary_dark="#f8fafc",
    block_background_fill="#ffffff",
    block_background_fill_dark="#ffffff",
    input_background_fill="#ffffff",
    input_background_fill_dark="#ffffff",
    body_text_color="#1e1b4b",
    body_text_color_dark="#1e1b4b",
)


# ---------------------------------------------------------
# Pipeline
# ---------------------------------------------------------


@lru_cache(maxsize=1)
def get_pipeline() -> DataPipeline:
    """Instancie le pipeline une seule fois avec OpenAI."""
    return DataPipeline(llm_client=OpenAIClient())


# ---------------------------------------------------------
# Détection des colonnes et formatage
# ---------------------------------------------------------


def _norm(name: Any) -> str:
    return str(name).lower()


def is_id_column(name: Any) -> bool:
    return _norm(name).endswith("id")


def is_time_column(name: Any) -> bool:
    return any(word in _norm(name) for word in TIME_WORDS)


def clean_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Convertit les colonnes Decimal/objets numériques en vrais nombres."""
    df = df.copy().reset_index(drop=True)

    for col in df.columns:
        if df[col].dtype != object:
            continue
        values = df[col].dropna()
        if len(values) and all(
            isinstance(v, (Decimal, int, float)) and not isinstance(v, bool)
            for v in values
        ):
            df[col] = pd.to_numeric(df[col], errors="coerce")

    return month_names(df)


def month_names(df: pd.DataFrame) -> pd.DataFrame:
    """Remplace un numéro de mois (8) par son nom (août), en catégorie ordonnée."""
    for col in df.columns:
        if _norm(col) not in MONTH_COLUMN_NAMES:
            continue

        values = pd.to_numeric(df[col], errors="coerce")

        if values.notna().all() and values.between(1, 12).all() and (values % 1 == 0).all():
            df[col] = pd.Categorical(
                [MONTHS_FR[int(v) - 1] for v in values],
                categories=MONTHS_FR,
                ordered=True,
            )

    return df


def measure_columns(df: pd.DataFrame) -> list:
    """Colonnes d'indicateurs chiffrés (hors identifiants et dates)."""
    return [
        col
        for col in df.select_dtypes(include="number").columns
        if not is_id_column(col) and not is_time_column(col)
    ]


def label_columns(df: pd.DataFrame) -> list:
    """Colonnes d'axes d'analyse (région, produit, mois...)."""
    measures = set(measure_columns(df))
    return [c for c in df.columns if c not in measures and not is_id_column(c)]


def pretty_column(name: Any) -> str:
    """Chiffre_Affaires_Total -> Chiffre d'affaires total."""
    text = str(name).replace("_", " ").strip().lower()
    text = re.sub(r"^nom ", "", text)
    for pattern, replacement in LABEL_REPLACEMENTS:
        text = re.sub(pattern, replacement, text)
    return text[:1].upper() + text[1:]


def detect_unit(name: Any) -> str:
    """Devine l'unité d'un indicateur d'après son nom."""
    n = _norm(name)

    if any(word in n for word in ("taux", "pct", "pourcent", "ratio", "%")):
        return "%"

    euro_words = ("chiffre", "affaires", "marge", "montant", "prix", "cout", "revenu", "benefice")
    if any(word in n for word in euro_words) or re.search(r"(^|_)ca($|_)", n):
        return "€"

    if "quantit" in n:
        return "unités"

    return ""


def format_number(value: int) -> str:
    return f"{value:,}".replace(",", "\u00a0")


def plural(count: int, word: str) -> str:
    return f"{format_number(count)} {word}{'s' if count > 1 else ''}"


def format_value(value: Any, unit: str = "", compact: bool = False) -> str:
    """Formate un nombre à la française : 13 136 952 € ou 13,1 M€."""
    number = float(value)

    if unit == "%" and abs(number) <= 1:
        number *= 100

    if compact and unit == "€" and abs(number) >= 1e6:
        scaled, suffix = (
            (number / 1e9, "Md€") if abs(number) >= 1e9 else (number / 1e6, "M€")
        )
        return f"{scaled:.1f}".replace(".", ",") + "\u00a0" + suffix

    if unit == "%":
        decimals = 0 if number.is_integer() else 1
    else:
        decimals = 0 if number.is_integer() or abs(number) >= 1000 else 2

    text = f"{number:,.{decimals}f}".replace(",", "\u00a0").replace(".", ",")
    return f"{text}\u00a0{unit}" if unit else text


def format_label(value: Any) -> str:
    """Affiche une valeur d'axe sans décimales parasites (2024, pas 2024.0)."""
    if isinstance(value, pd.Timestamp):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def is_additive(col: Any, values: pd.Series, unit: str) -> bool:
    """Une somme ou une part du total a-t-elle un sens pour cet indicateur ?"""
    return (
        values.min() >= 0
        and unit != "%"
        and not any(word in _norm(col) for word in NON_ADDITIVE_WORDS)
    )


# ---------------------------------------------------------
# Tableau : noms lisibles, nombres formatés, couleurs
# ---------------------------------------------------------


def _heat_style(ratio: float) -> str:
    r, g, b = (round(low + (high - low) * ratio) for low, high in zip(HEAT_LOW, HEAT_HIGH))
    text = "#ffffff" if ratio > 0.55 else TEXT_DARK
    return f"background-color: rgb({r},{g},{b}); color: {text};"


def readable_names(df: pd.DataFrame) -> list[str]:
    names = [pretty_column(c) for c in df.columns]
    return names if len(set(names)) == len(names) else [str(c) for c in df.columns]


def style_dataframe(df: pd.DataFrame):
    """Tableau lisible : noms propres, nombres formatés, lignes alternées, barres de données."""
    try:
        data = df.reset_index(drop=True)
        measures = set(measure_columns(data))
        names = readable_names(data)

        display = pd.DataFrame(index=data.index)
        for col, name in zip(data.columns, names):
            unit = detect_unit(col) if col in measures else ""
            display[name] = [
                ""
                if pd.isna(v)
                else (format_value(v, unit) if col in measures else format_label(v))
                for v in data[col]
            ]

        # Barre proportionnelle sur le premier indicateur
        bar_col = None
        first_measure = next((c for c in data.columns if c in measures), None)
        if first_measure is not None and len(data) > 1:
            top = data[first_measure].max()
            if pd.notna(top) and top > 0:
                bar_col = "Poids relatif"
                display[bar_col] = [
                    "" if pd.isna(v) else "█" * max(1, round(14 * max(v, 0) / top))
                    for v in data[first_measure]
                ]

        if len(data) > MAX_STYLED_ROWS:
            return display

        styles = pd.DataFrame("", index=display.index, columns=display.columns)
        for i in range(len(data)):
            styles.iloc[i, :] = f"background-color: {ZEBRA[i % 2]}; color: {TEXT_DARK};"

        for name, col in zip(names, data.columns):
            if col in measures:
                for i in range(len(data)):
                    styles.at[i, name] += " text-align: right;"

        if bar_col:
            for i in range(len(data)):
                styles.at[i, bar_col] = (
                    f"background-color: {ZEBRA[i % 2]}; color: #6366f1; letter-spacing: -1px;"
                )

        return display.style.apply(lambda _: styles, axis=None)

    except Exception:
        logger.exception("Mise en forme du tableau impossible")
        return df


# ---------------------------------------------------------
# Exports
# ---------------------------------------------------------


def export_frame(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.columns = readable_names(df)
    return out


def export_csv(df: pd.DataFrame) -> str:
    path = Path(tempfile.mkdtemp()) / "resultats.csv"
    export_frame(df).to_csv(path, index=False, encoding="utf-8-sig", sep=";", decimal=",")
    return str(path)


def export_excel(df: pd.DataFrame) -> str | None:
    """Export Excel mis en forme. Retourne None si openpyxl est absent."""
    try:
        from openpyxl.styles import Font, PatternFill

        path = Path(tempfile.mkdtemp()) / "resultats.xlsx"
        out = export_frame(df)

        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            out.to_excel(writer, index=False, sheet_name="Résultats")
            sheet = writer.sheets["Résultats"]

            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="4F46E5")

            for idx, col in enumerate(out.columns, start=1):
                longest = max([len(str(col))] + [len(str(v)) for v in out[col].head(200)])
                letter = sheet.cell(row=1, column=idx).column_letter
                sheet.column_dimensions[letter].width = min(longest + 3, 50)

        return str(path)

    except Exception:
        logger.warning("Export Excel indisponible (pip install openpyxl)", exc_info=True)
        return None


# ---------------------------------------------------------
# Réponse de repli (pipeline classique, sans agent)
# ---------------------------------------------------------


def analysis_label_column(question: str, df: pd.DataFrame) -> Any | None:
    """Détermine la dimension réellement analysée à partir de la question."""
    labels = label_columns(df)
    if not labels:
        return None

    q = _norm(question)

    if "mois" in q or "month" in q:
        for col in labels:
            if "nom_mois" in _norm(col):
                return col
        for col in labels:
            if _norm(col) in ("mois", "month"):
                return col

    if "région" in q or "region" in q:
        for col in labels:
            if "region" in _norm(col):
                return col

    if "produit" in q:
        for col in labels:
            if "produit" in _norm(col):
                return col

    if "année" in q or "annee" in q or "year" in q:
        for col in labels:
            n = _norm(col)
            if "annee" in n or "year" in n:
                return col

    return labels[0]


def build_answer(question: str, df: pd.DataFrame) -> str:
    """Phrase de réponse directe avec le chiffre clé."""
    measures = measure_columns(df)
    data = df.dropna(subset=measures[:1]) if measures else df
    if not measures or data.empty:
        return ""

    labels = label_columns(df)
    value_col = measures[0]
    analysis_label = analysis_label_column(question, df)
    unit = detect_unit(value_col)
    metric = pretty_column(value_col).lower()

    q = question.lower()
    wants_min = any(w in q for w in ("moins", "minimum", "faible", "pire", "dernier", "plus bas"))
    wants_max = not wants_min and any(
        w in q for w in ("plus", "meilleur", "maximum", "premier", "top", "élevé")
    )

    ranked = data.sort_values(value_col, ascending=False)
    row = ranked.iloc[-1] if wants_min else ranked.iloc[0]
    value = f"**{format_value(row[value_col], unit, compact=True)}**"

    if not labels:
        return f"{pretty_column(value_col)} : {value}."

    if analysis_label is not None:
        label = f"**{format_label(row[analysis_label])}**"

        if is_time_column(analysis_label) and len(data) > 1:
            phrase = "Point bas" if wants_min else "Pic"
            return f"{phrase} en {label} — {metric} : {value}."

        if wants_min:
            return f"{label} affiche la valeur la plus faible — {metric} : {value}."

        if wants_max or len(data) > 1:
            return f"{label} affiche la valeur la plus élevée — {metric} : {value}."

        return f"{label} — {metric} : {value}."

    return f"{metric} : {value}."


def build_insights(question: str, df: pd.DataFrame) -> list[str]:
    """Synthèse métier concise (repli sans agent)."""
    measures = measure_columns(df)
    if not measures:
        return []

    data = df.dropna(subset=[measures[0]]).copy()
    if len(data) < 2:
        return []

    value_col = measures[0]
    unit = detect_unit(value_col)
    analysis_label = analysis_label_column(question, data)
    if analysis_label is None:
        return []

    ranked = data.sort_values(value_col, ascending=False)
    best, last = ranked.iloc[0], ranked.iloc[-1]
    values = ranked[value_col]

    def fmt(value: Any) -> str:
        return format_value(value, unit, compact=True)

    dimension = pretty_column(analysis_label).lower()

    lines = [
        f"**{dimension.capitalize()} le plus élevé** : "
        f"**{format_label(best[analysis_label])}** — **{fmt(best[value_col])}**.",
        f"**{dimension.capitalize()} le plus faible** : "
        f"**{format_label(last[analysis_label])}** — **{fmt(last[value_col])}**.",
    ]

    if is_additive(value_col, values, unit):
        lines.append(f"**Total** : **{fmt(values.sum())}** · **moyenne** : **{fmt(values.mean())}**.")
    else:
        lines.append(f"**Moyenne** : **{fmt(values.mean())}**.")

    return lines


# ---------------------------------------------------------
# Panneau de droite : KPI, traçabilité, suggestions
# ---------------------------------------------------------


def metric_card(label: str, value: str, sub: str = "", alt: bool = False) -> str:
    sub_html = f'<div class="metric-sub">{sub}</div>' if sub else ""
    css = "metric-card alt" if alt else "metric-card"
    return (
        f'<div class="{css}"><div class="metric-label">{html.escape(label)}</div>'
        f'<div class="metric-value">{value}</div>{sub_html}</div>'
    )


def format_last_updated(value: Any) -> str | None:
    if not value:
        return None
    try:
        ts = pd.Timestamp(value)
        return f"{ts.day:02d}/{ts.month:02d}/{ts.year} à {ts.hour:02d}:{ts.minute:02d}"
    except Exception:
        return str(value)


def delta_html(pct: float | None, label: str = "") -> str:
    """▲ +4,2 % vs 2023 (vert) / ▼ -3,1 % (rouge)."""
    if pct is None:
        return ""
    if abs(pct) < 0.0005:
        arrow, css = "▬", "flat"
    elif pct > 0:
        arrow, css = "▲", "up"
    else:
        arrow, css = "▼", "down"
    text = f"{pct * 100:+.1f} %".replace(".", ",")
    return (
        f'<span class="delta {css}">{arrow} {text}</span> '
        f'<span class="delta-label">{html.escape(label)}</span>'
    )


def build_kpi_html(
    df: pd.DataFrame,
    elapsed: float | None,
    question: str = "",
    last_updated: Any = None,
    compare: dict | None = None,
    data_until: str | None = None,
) -> str:
    """En-tête : pastilles d'état uniquement (les chiffres clés sont dans « À retenir »)."""
    badges = [
        f'<span class="kpi-badge">{plural(len(df), "résultat")}</span>',
        '<span class="kpi-badge green">🟢 SQL connecté</span>',
        '<span class="kpi-badge green">🔒 Requête validée · lecture seule</span>',
    ]

    if elapsed is not None:
        seconds = f"{elapsed:.1f}".replace(".", ",")
        badges.append(f'<span class="kpi-badge">⏱ {seconds} s</span>')

    if data_until:
        badges.append(
            f'<span class="kpi-badge">📅 Données jusqu\'à {html.escape(str(data_until))}</span>'
        )

    last_updated_text = format_last_updated(last_updated)
    if last_updated_text:
        badges.append(
            f'<span class="kpi-badge">↻ Mise à jour : {html.escape(last_updated_text)}</span>'
        )

    return f'<div class="kpi-section"><div class="kpi-status-row">{"".join(badges)}</div></div>'


def build_trace(result, elapsed: float | None) -> str:
    """Traçabilité : d'où vient le chiffre."""
    parts: list[str] = []

    explanation = getattr(result, "explanation", None)
    if explanation:
        parts.append(f"**Méthode** : {explanation}")

    tables = getattr(result, "tables_used", None)
    if tables:
        parts.append("**Tables utilisées** : " + ", ".join(f"`{t}`" for t in tables))

    columns = getattr(result, "columns_used", None)
    if columns:
        parts.append("**Colonnes utilisées** : " + ", ".join(f"`{c}`" for c in columns))

    assumptions = getattr(result, "assumptions", None)
    if assumptions:
        parts.append("**Hypothèses**\n" + "\n".join(f"- {a}" for a in assumptions))

    parts.append("**Contrôle de sécurité** : requêtes validées avant exécution (lecture seule).")

    if elapsed is not None:
        parts.append(f"**Temps de réponse** : {f'{elapsed:.1f}'.replace('.', ',')} s")

    return "\n\n".join(parts)


# def build_suggestions(df: pd.DataFrame) -> list[list[str]]:
#     """Questions de suivi par défaut (si l'assistant n'en propose pas)."""
#     measures = measure_columns(df)
#     if not measures:
#         return []

#     metric = re.sub(r" totale?$", "", pretty_column(measures[0]).lower())
#     metric = metric[:1].upper() + metric[1:]

#     labels = label_columns(df)
#     current = _norm(labels[0]) if labels else ""

#     suggestions = [
#         f"{metric} par {word}"
#         for key, word in FOLLOW_UP_DIMENSIONS
#         if key not in current
#     ]
#     suggestions.append(f"Prévois {metric.lower()} pour les 6 prochains mois")

#     return [[s] for s in suggestions[:4]]


# ---------------------------------------------------------
# Graphique générique piloté par l'assistant
# ---------------------------------------------------------


def _resolve_column(df: pd.DataFrame, name: Any):
    if not name:
        return None
    lookup = {str(c).lower(): c for c in df.columns}
    return lookup.get(str(name).strip().lower())


def _hover_format(unit: str, values: list[float]) -> str:
    """Format Plotly du survol, accolade fermante incluse."""
    small = bool(values) and max(abs(v) for v in values) < 100
    pattern = ",.2f" if small else ",.0f"
    if unit == "%":
        return ",.1f} %"
    if unit == "€":
        return f"{pattern}}} €"
    if unit == "unités":
        return f"{pattern}}} unités"
    return f"{pattern}}}"


def _is_orderable(series: pd.Series) -> bool:
    """Colonne dont l'ordre naturel est fiable (dates, nombres, catégories ordonnées)."""
    return (
        pd.api.types.is_datetime64_any_dtype(series)
        or pd.api.types.is_numeric_dtype(series)
        or isinstance(series.dtype, pd.CategoricalDtype)
    )


_YEAR_RE = re.compile(r"\b(?:19|20)\d{2}\b")

############################################################""
GEOJSON_PATH = BASE_DIR / "data" / "regions.geojson"


def _plain_key(text: Any) -> str:
    import unicodedata

    t = unicodedata.normalize("NFD", str(text).lower().replace("’", "'"))
    return "".join(c for c in t if unicodedata.category(c) != "Mn").strip()


@lru_cache(maxsize=1)
def _load_regions_geojson():
    """Charge le GeoJSON des régions (None si le fichier est absent)."""
    if not GEOJSON_PATH.is_file():
        return None, {}
    import json

    gj = json.loads(GEOJSON_PATH.read_text(encoding="utf-8"))
    names = {
        _plain_key(f["properties"]["nom"]): f["properties"]["nom"]
        for f in gj.get("features", [])
        if "nom" in (f.get("properties") or {})
    }
    return gj, names

############################""""

def _region_bbox(gj, feature_name: str):
    """(lon_min, lon_max, lat_min, lat_max) d'une région du GeoJSON."""
    for f in gj.get("features", []):
        if (f.get("properties") or {}).get("nom") != feature_name:
            continue
        pts: list = []

        def walk(c):
            if c and isinstance(c[0], (int, float)):
                pts.append(c)
            else:
                for s in c:
                    walk(s)

        walk(f["geometry"]["coordinates"])
        lons, lats = [p[0] for p in pts], [p[1] for p in pts]
        pad = max(max(lons) - min(lons), max(lats) - min(lats)) * 0.25
        return min(lons) - pad, max(lons) + pad, min(lats) - pad, max(lats) + pad
    return None


def _france_map(data, x_col, value_col, title, hover_fmt, pretty_label, selected=None):
    gj, names = _load_regions_geojson()
    if not gj:
        return None

    rows = [
        (names.get(_plain_key(r)), float(v))
        for r, v in zip(data[x_col], data[value_col])
    ]
    rows = [(n, v) for n, v in rows if n]
    if len(rows) < max(3, int(len(data) * 0.6)):
        return None   # trop de noms non reconnus : on garde les barres

    hover = (
        f"<b>%{{location}}</b><br>{pretty_label} : "
        f"<b>%{{z:{hover_fmt}</b><extra></extra>"
    )

    fig = go.Figure(
        go.Choropleth(
            geojson=gj,
            locations=[n for n, _ in rows],
            z=[v for _, v in rows],
            featureidkey="properties.nom",
            colorscale=[[0, "#E0E7FF"], [1, "#4F46E5"]],
            marker_line_color="white",
            marker_line_width=0.8,
            colorbar=dict(title=None, tickformat="~s"),
            hovertemplate=hover,
        )
    )

    sel_name = names.get(_plain_key(selected)) if selected else None
    bbox = _region_bbox(gj, sel_name) if sel_name else None

    if bbox:
        value = next(v for n, v in rows if n == sel_name) if any(n == sel_name for n, _ in rows) else None
        if value is not None:
            fig.add_trace(
                go.Choropleth(
                    geojson=gj,
                    locations=[sel_name],
                    z=[value],
                    featureidkey="properties.nom",
                    colorscale=[[0, "#F59E0B"], [1, "#F59E0B"]],
                    showscale=False,
                    marker_line_color="#111827",
                    marker_line_width=2.5,
                    hovertemplate=hover,
                )
            )
        fig.update_geos(
            visible=False,
            projection_type="mercator",
            lonaxis_range=[bbox[0], bbox[1]],
            lataxis_range=[bbox[2], bbox[3]],
        )
    else:
        fig.update_geos(fitbounds="locations", visible=False)

    fig.update_layout(
        title=dict(text=title, x=0, xanchor="left", font=dict(size=18, color="#27245C")),
        margin=dict(l=0, r=0, t=70, b=0),
        height=340,
        paper_bgcolor="white",
        geo=dict(bgcolor="white"),
    )
    return fig


def _focus_map(rdf: pd.DataFrame, region: str, years: list[str]):
    """Carte de toutes les régions, recadrée sur la région choisie."""
    labels, measures = label_columns(rdf), measure_columns(rdf)
    x = next((c for c in labels if "region" in _norm(c)), None)
    if x is None or not measures or not rdf[x].is_unique:
        return None
    y = measures[0]
    hover = _hover_format(detect_unit(y), [float(v) for v in rdf[y]])
    title = f"{pretty_column(y)} par région{_period_suffix(years)} — {region}"
    return _france_map(rdf, x, y, title, hover, pretty_column(y), selected=region)


def _rank_chip(rdf: pd.DataFrame, region: str):
    """Rang et part de la région dans le total (calculés en code)."""
    labels, measures = label_columns(rdf), measure_columns(rdf)
    x = next((c for c in labels if "region" in _norm(c)), None)
    if x is None or not measures:
        return None
    y = measures[0]
    unit = detect_unit(y)
    if not is_additive(y, rdf[y].dropna(), unit):
        return None
    ranked = rdf.dropna(subset=[y]).sort_values(y, ascending=False).reset_index(drop=True)
    hit = ranked.index[ranked[x].map(_plain_key) == _plain_key(region)]
    if len(hit) == 0:
        return None
    i = int(hit[0])
    share = float(ranked.loc[i, y]) / float(ranked[y].sum()) * 100
    return (
        "",
        f"Rang {i + 1} / {len(ranked)}",
        f"{region} · {format_value(ranked.loc[i, y], unit, compact=True)} · "
        f"{share:.1f} % du total".replace(".", ","),
    )


def _extreme_colors(values: list[float], base: str = "#4F46E5"):
    """Maximum en vert, minimum en rouge, le reste en couleur de base."""
    if len(values) < 3 or max(values) == min(values):
        return base
    hi, lo = max(values), min(values)
    return ["#16A34A" if v == hi else "#EF4444" if v == lo else base for v in values]


def _ordered_points(df: pd.DataFrame, value_col) -> list[tuple[str, float]] | None:
    """Série chronologique [(libellé, valeur)] si le résultat a un axe temps fiable."""
    labels = label_columns(df)
    year_col = next((c for c in labels if any(w in _norm(c) for w in ("annee", "year"))), None)
    month_col = next((c for c in labels if _norm(c) in MONTH_COLUMN_NAMES), None)
    date_col = next(
        (c for c in labels if any(w in _norm(c) for w in ("date", "jour", "periode", "période"))),
        None,
    )
    x = month_col or date_col or year_col
    if x is None:
        return None

    data = df.dropna(subset=[value_col])
    if month_col and x == month_col and year_col and data[year_col].nunique() > 1:
        data = data.sort_values([year_col, month_col])
        names = [f"{format_label(m)} {format_label(y)}" for m, y in zip(data[month_col], data[year_col])]
    elif _is_orderable(data[x]):
        data = data.sort_values(x)
        names = [format_label(v) for v in data[x]]
    else:
        return None
    return list(zip(names, data[value_col].astype(float)))


def make_chart_from_spec(
    df: pd.DataFrame,
    spec: dict | None = None,
    anomalies: list | None = None,
):
    """Graphique Plotly piloté par l'assistant (type, x, y, years), avec repli automatique."""
    spec = spec or {}
    kind = str(spec.get("type") or "auto").lower()
    requested = kind

    if kind == "none" or len(df) < 2:
        return None

    measures, labels = measure_columns(df), label_columns(df)

    y_cols = [
        c
        for c in (_resolve_column(df, p) for p in str(spec.get("y") or "").split(","))
        if c is not None and c in measures
    ] or measures[:1]

    if not y_cols or not labels:
        return None

    year_col = next((c for c in labels if any(w in _norm(c) for w in ("annee", "year"))), None)
    month_col = next((c for c in labels if _norm(c) in MONTH_COLUMN_NAMES), None)
    date_col = next(
        (c for c in labels if any(w in _norm(c) for w in ("date", "jour", "periode", "période"))),
        None,
    )

    x_col = _resolve_column(df, spec.get("x")) or month_col or date_col or year_col or labels[0]
    if x_col in y_cols:
        return None

    data = df.dropna(subset=y_cols[:1]).copy()
    time_like = is_time_column(x_col) or pd.api.types.is_datetime64_any_dtype(data[x_col])

    # ---- Libellés de l'axe X ----
    if month_col and x_col == month_col and year_col and data[year_col].nunique() > 1:
        data = data.sort_values([year_col, month_col])
        names = [
            f"{format_label(m)} {format_label(y)}"
            for m, y in zip(data[month_col], data[year_col])
        ]
    else:
        if time_like and _is_orderable(data[x_col]):
            data = data.sort_values(x_col)
        names = [format_label(v) for v in data[x_col]]

    # ---- Année(s) pour le titre ----
    raw_years = spec.get("years") or []
    if isinstance(raw_years, (int, str)):
        raw_years = [raw_years]
    years = list(dict.fromkeys(str(y).strip() for y in raw_years if str(y).strip().isdigit()))

    if not years and year_col and year_col in data.columns:
        years = sorted({format_label(v) for v in data[year_col].dropna()})

    year_already_visible = x_col == year_col or any(_YEAR_RE.search(n) for n in names)
    period_title = "" if year_already_visible else _period_suffix(years)

    first = y_cols[0]
    unit = detect_unit(first)
    hover_fmt = _hover_format(unit, [float(v) for v in data[first]])

    title = (
        f"{', '.join(pretty_column(c) for c in y_cols)} "
        f"par {pretty_column(x_col).lower()}"
        f"{period_title}"
    )
    x_title = ""

    # ---- Carte de France : axe région, une seule valeur par région ----
    if (
        requested in ("auto", "map", "line")
        and "region" in _norm(x_col)
        and len(y_cols) == 1
        and data[x_col].is_unique
    ):
        map_fig = _france_map(data, x_col, first, title, hover_fmt, pretty_column(first))
        if map_fig is not None:
            return map_fig

    # Une courbe n'a de sens que sur un axe temporel : sinon, barres
    if kind in ("auto", "map") or (kind == "line" and not time_like):
        kind = "line" if time_like else "bar"

    palette = ["#4F46E5", "#7C3AED", "#0EA5E9", "#F59E0B", "#10B981"]
    fig = go.Figure()

    # ---- PIE ----
    if kind == "pie":
        pie = data.sort_values(first, ascending=False).head(8)
        fig.add_trace(
            go.Pie(
                labels=[format_label(v) for v in pie[x_col]],
                values=pie[first],
                hole=0.4,
                textinfo="label+percent",
            )
        )

    # ---- BAR : valeurs affichées, max en vert, min en rouge ----
    elif kind == "bar":
        if len(data) > 12 and len(y_cols) == 1:
            top = data.sort_values(first, ascending=False).head(15).iloc[::-1]
            vals = [float(v) for v in top[first]]
            fig.add_trace(
                go.Bar(
                    x=vals,
                    y=[format_label(v)[:40] for v in top[x_col]],
                    orientation="h",
                    marker_color=_extreme_colors(vals, palette[0]),
                    text=[format_value(v, unit, compact=True) for v in vals],
                    textposition="outside",
                    cliponaxis=False,
                    hovertemplate=(
                        f"<b>%{{y}}</b><br>{pretty_column(first)} : "
                        f"<b>%{{x:{hover_fmt}</b><extra></extra>"
                    ),
                )
            )
        else:
            for i, col in enumerate(y_cols):
                vals = [float(v) for v in data[col]]
                col_unit = detect_unit(col)
                fig.add_trace(
                    go.Bar(
                        x=names,
                        y=vals,
                        name=pretty_column(col),
                        marker_color=(
                            _extreme_colors(vals, palette[0])
                            if len(y_cols) == 1
                            else palette[i % len(palette)]
                        ),
                        text=(
                            [format_value(v, col_unit, compact=True) for v in vals]
                            if len(vals) <= 15 else None
                        ),
                        textposition="outside",
                        cliponaxis=False,
                        hovertemplate=(
                            f"<b>%{{x}}</b><br>{pretty_column(col)} : "
                            f"<b>%{{y:{hover_fmt}</b><extra></extra>"
                        ),
                    )
                )

    # ---- LINE : points max/min colorés ----
    else:
        for i, col in enumerate(y_cols):
            color = palette[i % len(palette)]
            vals = [float(v) for v in data[col]]
            fig.add_trace(
                go.Scatter(
                    x=names,
                    y=vals,
                    name=pretty_column(col),
                    mode="lines+markers",
                    line=dict(color=color, width=3),
                    marker=dict(
                        size=8,
                        color=_extreme_colors(vals, color) if len(y_cols) == 1 else color,
                    ),
                    fill="tozeroy" if len(y_cols) == 1 else None,
                    fillcolor="rgba(79,70,229,0.12)",
                    hovertemplate=(
                        f"<b>%{{x}}</b><br>{pretty_column(col)} : "
                        f"<b>%{{y:{hover_fmt}</b><extra></extra>"
                    ),
                )
            )

        matches = [a for a in (anomalies or []) if a.get("periode") in names]
        if matches and len(y_cols) == 1:
            fig.add_trace(
                go.Scatter(
                    x=[a["periode"] for a in matches],
                    y=[a["valeur"] for a in matches],
                    mode="markers",
                    name="Anomalie",
                    marker=dict(size=15, color="rgba(0,0,0,0)", line=dict(color="#dc2626", width=3)),
                    hovertemplate=(
                        f"<b>%{{x}}</b><br>Anomalie : <b>%{{y:{hover_fmt}</b><extra></extra>"
                    ),
                )
            )

    fig.update_layout(
        title=dict(text=title, x=0, xanchor="left", font=dict(size=18, color="#27245C")),
        xaxis=dict(title=dict(text=x_title, font=dict(size=12)), type="category", tickangle=-45),
        yaxis=dict(title=None, tickformat=","),
        hovermode="closest" if kind in ("bar", "pie") else "x unified",
        margin=dict(l=60, r=30, t=70, b=90),
        height=340,
        plot_bgcolor="white",
        paper_bgcolor="white",
        showlegend=len(y_cols) > 1 or bool(anomalies),
        barmode="group",
    )
    fig.update_xaxes(showgrid=False)
    fig.update_yaxes(showgrid=True, gridcolor="#E5E7EB")
    return fig


############################################################""

def _period_suffix(years: list[str]) -> str:
    """['2024'] -> ' — 2024' ; ['2023','2024'] -> ' — 2023, 2024' ; 2022..2024 -> ' — 2022–2024'."""
    nums = sorted({int(y) for y in years})
    if not nums:
        return ""
    if len(nums) > 2 and nums[-1] - nums[0] == len(nums) - 1:
        return f" — {nums[0]}–{nums[-1]}"
    return " — " + ", ".join(str(n) for n in nums)


def _years_from(*texts: str | None) -> list[str]:
    """Années explicites : le premier texte qui en contient gagne (SQL avant question)."""
    for text in texts:
        found = sorted(set(_YEAR_RE.findall(text or "")))
        if found:
            return found
    return []

# def make_chart_from_spec(
#     df: pd.DataFrame,
#     spec: dict | None = None,
#     anomalies: list | None = None,
# ):
#     """Graphique Plotly piloté par l'assistant (type, x, y, years), avec repli automatique."""
#     spec = spec or {}
#     kind = str(spec.get("type") or "auto").lower()

#     if kind == "none" or len(df) < 2:
#         return None

#     measures, labels = measure_columns(df), label_columns(df)

#     y_cols = [
#         c
#         for c in (_resolve_column(df, p) for p in str(spec.get("y") or "").split(","))
#         if c is not None and c in measures
#     ] or measures[:1]

#     if not y_cols or not labels:
#         return None

#     year_col = next((c for c in labels if any(w in _norm(c) for w in ("annee", "year"))), None)
#     month_col = next((c for c in labels if _norm(c) in MONTH_COLUMN_NAMES), None)
#     date_col = next(
#         (c for c in labels if any(w in _norm(c) for w in ("date", "jour", "periode", "période"))),
#         None,
#     )

#     x_col = _resolve_column(df, spec.get("x")) or month_col or date_col or year_col or labels[0]
#     if x_col in y_cols:
#         return None

#     data = df.dropna(subset=y_cols[:1]).copy()
#     time_like = is_time_column(x_col) or pd.api.types.is_datetime64_any_dtype(data[x_col])

#     # ---- Libellés de l'axe X ----
#     if month_col and x_col == month_col and year_col and data[year_col].nunique() > 1:
#         # Mois sur plusieurs années : « mars 2024 », triés par (année, mois)
#         data = data.sort_values([year_col, month_col])
#         names = [
#             f"{format_label(m)} {format_label(y)}"
#             for m, y in zip(data[month_col], data[year_col])
#         ]
#     else:
#         # On ne trie que si l'ordre est fiable (pas de tri alphabétique de « janvier 2025 »)
#         if time_like and _is_orderable(data[x_col]):
#             data = data.sort_values(x_col)
#         names = [format_label(v) for v in data[x_col]]

#     # ---- Année(s) concernée(s) pour le titre ----
#     # 1) années trouvées dans le SQL ou la question (transmises par l'agent)
#     raw_years = spec.get("years") or []
#     if isinstance(raw_years, (int, str)):
#         raw_years = [raw_years]
#     years = list(dict.fromkeys(str(y).strip() for y in raw_years if str(y).strip().isdigit()))

#     # 2) sinon, années présentes dans une colonne du résultat
#     if not years and year_col and year_col in data.columns:
#         years = sorted({format_label(v) for v in data[year_col].dropna()})

#     # Inutile si l'année est déjà sur l'axe (colonne année ou libellés « mars 2025 »)
#     year_already_visible = x_col == year_col or any(_YEAR_RE.search(n) for n in names)
#     period_title = "" if year_already_visible else _period_suffix(years)

#     if kind == "auto":
#         kind = "line" if time_like else "bar"

#     first = y_cols[0]
#     unit = detect_unit(first)
#     hover_fmt = _hover_format(unit, [float(v) for v in data[first]])

#     title = (
#         f"{', '.join(pretty_column(c) for c in y_cols)} "
#         f"par {pretty_column(x_col).lower()}"
#         f"{period_title}"
#     )
#     x_title = ""

#     if (
#         month_col
#         and x_col == month_col
#         and year_col
#         and data[year_col].nunique() > 1
#     ):
#         # Mois sur plusieurs années :
#         # « mars 2024 », « avril 2024 », etc.
#         # triés par (année, mois)
#         data = data.sort_values([year_col, month_col])

#         names = [
#             f"{format_label(m)} {format_label(y)}"
#             for m, y in zip(data[month_col], data[year_col])
#         ]

#     else:
#         # On ne trie que si l'ordre est fiable.
#         # Les libellés texte comme « janvier 2025 »
#         # ne doivent pas être triés alphabétiquement.
#         if time_like and _is_orderable(data[x_col]):
#             data = data.sort_values(x_col)

#         names = [format_label(v) for v in data[x_col]]

#         # Axe X = mois et une seule année :
#         # le contexte temporel est déjà placé dans le titre.
#         if (
#             month_col
#             and x_col == month_col
#             and year_col
#             and data[year_col].nunique() == 1
#         ):
#             period_title = f" — {format_label(data[year_col].iloc[0])}"

#     if kind == "auto":
#         kind = "line" if time_like else "bar"

#     first = y_cols[0]
#     unit = detect_unit(first)

#     hover_fmt = _hover_format(
#         unit,
#         [float(v) for v in data[first]],
#     )

#     title = (
#         f"{', '.join(pretty_column(c) for c in y_cols)} "
#         f"par {pretty_column(x_col).lower()}"
#         f"{period_title}"
#     )

#     palette = [
#         "#4F46E5",
#         "#7C3AED",
#         "#0EA5E9",
#         "#F59E0B",
#         "#10B981",
#     ]

#     fig = go.Figure()

#     # ------------------------------------------------------------------
#     # PIE
#     # ------------------------------------------------------------------
#     if kind == "pie":
#         pie = data.sort_values(
#             first,
#             ascending=False,
#         ).head(8)

#         fig.add_trace(
#             go.Pie(
#                 labels=[format_label(v) for v in pie[x_col]],
#                 values=pie[first],
#                 hole=0.4,
#                 textinfo="label+percent",
#             )
#         )

#     # ------------------------------------------------------------------
#     # BAR
#     # ------------------------------------------------------------------
#     elif kind == "bar":

#         # Beaucoup de catégories :
#         # graphique horizontal avec Top 15.
#         if len(data) > 12 and len(y_cols) == 1:
#             top = (
#                 data
#                 .sort_values(first, ascending=False)
#                 .head(15)
#                 .iloc[::-1]
#             )

#             fig.add_trace(
#                 go.Bar(
#                     x=top[first],
#                     y=[
#                         format_label(v)[:40]
#                         for v in top[x_col]
#                     ],
#                     orientation="h",
#                     marker_color=palette[0],
#                     hovertemplate=(
#                         f"<b>%{{y}}</b><br>"
#                         f"{pretty_column(first)} : "
#                         f"<b>%{{x:{hover_fmt}</b>"
#                         f"<extra></extra>"
#                     ),
#                 )
#             )

#         else:
#             for i, col in enumerate(y_cols):
#                 fig.add_trace(
#                     go.Bar(
#                         x=names,
#                         y=data[col],
#                         name=pretty_column(col),
#                         marker_color=palette[i % len(palette)],
#                         hovertemplate=(
#                             f"<b>%{{x}}</b><br>"
#                             f"{pretty_column(col)} : "
#                             f"<b>%{{y:{hover_fmt}</b>"
#                             f"<extra></extra>"
#                         ),
#                     )
#                 )

#     # ------------------------------------------------------------------
#     # LINE
#     # ------------------------------------------------------------------
#     else:
#         for i, col in enumerate(y_cols):
#             color = palette[i % len(palette)]

#             fig.add_trace(
#                 go.Scatter(
#                     x=names,
#                     y=data[col],
#                     name=pretty_column(col),
#                     mode="lines+markers",
#                     line=dict(
#                         color=color,
#                         width=3,
#                     ),
#                     marker=dict(
#                         size=8,
#                         color=color,
#                     ),
#                     fill="tozeroy" if len(y_cols) == 1 else None,
#                     fillcolor="rgba(79,70,229,0.12)",
#                     hovertemplate=(
#                         f"<b>%{{x}}</b><br>"
#                         f"{pretty_column(col)} : "
#                         f"<b>%{{y:{hover_fmt}</b>"
#                         f"<extra></extra>"
#                     ),
#                 )
#             )

#         # --------------------------------------------------------------
#         # Anomalies détectées :
#         # cercle rouge sur les points concernés
#         # --------------------------------------------------------------
#         matches = [
#             a
#             for a in (anomalies or [])
#             if a.get("periode") in names
#         ]

#         if matches and len(y_cols) == 1:
#             fig.add_trace(
#                 go.Scatter(
#                     x=[a["periode"] for a in matches],
#                     y=[a["valeur"] for a in matches],
#                     mode="markers",
#                     name="Anomalie",
#                     marker=dict(
#                         size=15,
#                         color="rgba(0,0,0,0)",
#                         line=dict(
#                             color="#dc2626",
#                             width=3,
#                         ),
#                     ),
#                     hovertemplate=(
#                         f"<b>%{{x}}</b><br>"
#                         f"Anomalie : "
#                         f"<b>%{{y:{hover_fmt}</b>"
#                         f"<extra></extra>"
#                     ),
#                 )
#             )

#     # ------------------------------------------------------------------
#     # MISE EN FORME
#     # ------------------------------------------------------------------
#     fig.update_layout(
#         title=dict(
#             text=title,
#             x=0,
#             xanchor="left",
#             font=dict(
#                 size=18,
#                 color="#27245C",
#             ),
#         ),
#         xaxis=dict(
#             title=dict(
#                 text=x_title,
#                 font=dict(size=12),
#             ),
#             type="category",
#             tickangle=-45,
#         ),
#         yaxis=dict(
#             title=None,
#             tickformat=",",
#         ),
#         hovermode=(
#             "closest"
#             if kind in ("bar", "pie")
#             else "x unified"
#         ),
#         margin=dict(
#             l=60,
#             r=30,
#             t=70,
#             b=90,
#         ),
#         height=430,
#         plot_bgcolor="white",
#         paper_bgcolor="white",
#         showlegend=(
#             len(y_cols) > 1
#             or bool(anomalies)
#         ),
#         barmode="group",
#     )

#     fig.update_xaxes(showgrid=False)
#     fig.update_yaxes(
#         showgrid=True,
#         gridcolor="#E5E7EB",
#     )

#     return fig


def make_forecast_chart(fr) -> tuple[Any, str] | None:
    """Graphique de prévision : historique, prévision, zone d'incertitude."""
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.ticker import FuncFormatter
    except ImportError:
        return None

    try:
        date_col, value_col = fr.date_column, fr.value_column

        history = fr.historical[[date_col, value_col]].copy()
        forecast = fr.forecast.copy()
        history[date_col] = pd.to_datetime(history[date_col])
        forecast[date_col] = pd.to_datetime(forecast[date_col])

        unit = detect_unit(value_col)

        fig, ax = plt.subplots(figsize=(11, 5.8), dpi=150)
        ax.yaxis.set_major_formatter(
            FuncFormatter(lambda v, _: format_value(v, unit, compact=True))
        )

        ax.plot(
            history[date_col], history[value_col],
            color=INDIGO, linewidth=2.4, marker="o", markersize=4, label="Historique",
        )
        ax.plot(
            forecast[date_col], forecast[value_col],
            color="#7c3aed", linewidth=2.8, marker="o", markersize=4,
            linestyle="--", label="Prévision",
        )

        if {"Lower", "Upper"} <= set(forecast.columns):
            ax.fill_between(
                forecast[date_col], forecast["Lower"], forecast["Upper"],
                color="#8b5cf6", alpha=0.16, label="Zone d'incertitude",
            )

        if not history.empty and not forecast.empty:
            last_history = history[date_col].max()
            ax.axvline(last_history, color="#94a3b8", linestyle=":", linewidth=1.2)
            ax.text(
                last_history, ax.get_ylim()[1], "  Début prévision",
                va="top", fontsize=9, color="#64748b",
            )

        ax.set_title(
            f"{pretty_column(value_col)} — historique et prévision",
            loc="left", fontsize=14, fontweight="bold", color=TEXT_DARK, pad=14,
        )
        ax.set_xlabel("")
        ax.grid(axis="y", alpha=0.18)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)

        ax.legend(loc="upper left", frameon=False, ncol=3)
        fig.autofmt_xdate()
        fig.tight_layout()

        path = Path(tempfile.mkdtemp()) / "prevision.png"
        fig.savefig(path, bbox_inches="tight", facecolor="white")

        return fig, str(path)

    except Exception:
        logger.exception("Génération du graphique de prévision impossible")
        return None


# ---------------------------------------------------------
# Prévision : tableau, KPI, panneau
# ---------------------------------------------------------

RELIABILITY_BADGES = {
    "bonne": ("green", "✅"),
    "moyenne": ("orange", "⚠️"),
    "faible": ("red", "⛔"),
}


def month_label(value: Any) -> str:
    """2026-10-01 -> octobre 2026."""
    ts = pd.Timestamp(value)
    return f"{MONTHS_FR[ts.month - 1]} {ts.year}"


def same_period_last_year(fr) -> float | None:
    """Total de l'historique sur les mêmes mois un an plus tôt."""
    history = fr.historical.copy()
    history[fr.date_column] = pd.to_datetime(history[fr.date_column])
    series = history.set_index(fr.date_column)[fr.value_column]

    previous = [
        pd.Timestamp(d) - pd.DateOffset(years=1)
        for d in pd.to_datetime(fr.forecast[fr.date_column])
    ]

    if not all(p in series.index for p in previous):
        return None

    total = float(series.loc[previous].sum())
    return total if total > 0 else None


def forecast_table(fr) -> pd.DataFrame:
    """Tableau métier : historique + prévisions."""
    date_col, value_col = fr.date_column, fr.value_column

    history = fr.historical[[date_col, value_col]].copy()
    history.columns = ["Date", value_col]
    history["Type"] = "Historique"

    projected = pd.DataFrame(
        {
            "Date": fr.forecast[date_col],
            value_col: fr.forecast[value_col],
            "Type": "Prévision",
        }
    )

    table = pd.concat([history, projected], ignore_index=True)
    first = ["Date", "Type"]
    return table[first + [c for c in table.columns if c not in first]]


def build_forecast_reply(
    result,
    panel: dict[str, Any],
    elapsed: float | None = None,
) -> tuple[str, dict[str, Any]]:
    """Réponse métier + panneau pour un résultat de prévision."""
    fr = result.forecast_result

    date_col, value_col = fr.date_column, fr.value_column
    unit = detect_unit(value_col)
    metric = pretty_column(value_col).lower()

    forecast = fr.forecast
    count = len(forecast)
    total = float(forecast[value_col].sum())

    previous = same_period_last_year(fr)
    variation = (total / previous - 1) * 100 if previous else None

    reliability = getattr(fr, "reliability", "faible")
    css, icon = RELIABILITY_BADGES.get(reliability, ("", "ℹ️"))
    metrics = fr.metrics or {}
    mape = metrics.get("MAPE")
    reason = getattr(fr, "selection_reason", "") or ""
    folds = getattr(fr, "backtest_folds", 0)
    window = getattr(fr, "backtest_horizon", 0)

    notes = list(
        dict.fromkeys(
            [
                *(getattr(fr, "quality_warnings", None) or []),
                *(getattr(result, "warnings", None) or []),
            ]
        )
    )

    if variation is None:
        variation_text = ""
    elif round(variation) == 0:
        variation_text = "stable"
    else:
        variation_text = f"{variation:+.0f} %"

    sub = f"{variation_text} vs mêmes mois un an plus tôt" if variation_text else ""
    mape_text = f"MAPE {mape:.1f} %".replace(".", ",") if mape is not None else "non validée"

    cards = metric_card(
        f"{pretty_column(value_col)} prévu · {count} mois",
        format_value(total, unit, compact=True),
        html.escape(sub),
    )
    cards += metric_card("Fiabilité", reliability.capitalize(), mape_text, alt=True)
    cards += metric_card(
        "Modèle retenu",
        f'<span style="font-size:1.05rem">{html.escape(fr.model_name)}</span>',
        alt=True,
    )

    badges = [
        '<span class="kpi-badge">📈 Prévision</span>',
        f'<span class="kpi-badge {css}">{icon} Fiabilité {reliability}</span>',
    ]
    if folds:
        badges.append(
            f'<span class="kpi-badge">🧪 Validée sur {folds} fenêtres de {window} mois</span>'
        )
    badges.append('<span class="kpi-badge green">🔒 Requête validée · lecture seule</span>')
    if elapsed is not None:
        seconds = f"{elapsed:.1f}".replace(".", ",")
        badges.append(f'<span class="kpi-badge">⏱ {seconds} s</span>')


###############################################""
    panel["kpi"] = (
        '<div class="kpi-section">'
        f'<div class="kpi-status-row">{"".join(badges)}</div>'
        "</div>"
    )

    forecast_chips = [
        (f"{pretty_column(value_col)} prévu · {count} mois",
         format_value(total, unit, compact=True) + (f" · {sub}" if sub else "")),
        ("Fiabilité", f"{reliability.capitalize()} · {mape_text}"),
        ("Modèle retenu", fr.model_name),
    ]
    panel["highlights"] = gr.update(
        value=(
            '<div class="hl-band"><span class="hl-label">À retenir</span>'
            + "".join(
                f'<div class="hl-chip"><span class="hl-title">{html.escape(t)}</span>'
                f'<span class="hl-text">{html.escape(x)}</span></div>'
                for t, x in forecast_chips
            )
            + "</div>"
        ),
        visible=True,
    )

    # ---- Message de repli (l'assistant rédige le sien quand il est disponible) ----
    dates = pd.to_datetime(forecast[date_col])
    period = (
        f"{month_label(dates.min())} à {month_label(dates.max())}"
        if count > 1
        else month_label(dates.iloc[0])
    )

    headline = f"**{format_value(total, unit, compact=True)}** prévus sur {count} mois"
    if variation_text == "stable":
        headline += ", soit un niveau **stable** par rapport aux mêmes mois un an plus tôt"
    elif variation_text:
        headline += f", soit **{variation_text}** par rapport aux mêmes mois un an plus tôt"

    peak = forecast.loc[forecast[value_col].idxmax()]
    hist_dates = pd.to_datetime(fr.historical[fr.date_column])

    lines = [
        f"📈 **Prévision — {metric}** · {period}",
        f"{headline}.",
        f"Mois le plus fort : **{month_label(peak[date_col])}** "
        f"({format_value(peak[value_col], unit, compact=True)}).",
        f"{icon} **Fiabilité {reliability}**",
    ]
    if reason:
        lines.append(f"🧠 {reason}")
    if reliability == "faible":
        lines.append("Prenez ces chiffres avec prudence.")
    lines.append(
        f"Historique utilisé : {month_label(hist_dates.min())} → "
        f"{month_label(hist_dates.max())} ({len(hist_dates)} mois)."
    )

    message = "\n\n".join(lines)
    if notes:
        message += "\n\n" + "\n".join(f"- ⚠️ {n}" for n in notes)
    message += "\n\n> Détail mois par mois dans le tableau, graphique dans **Visualisation**."

    # ---- Tableau, exports, graphique ----
    table = clean_dataframe(forecast_table(fr))

    panel["table"] = gr.update(value=style_dataframe(table), visible=True)
    panel["download"] = gr.update(value=export_csv(table), visible=True)

    excel_path = export_excel(table)
    panel["excel"] = gr.update(value=excel_path, visible=bool(excel_path))

    chart_result = make_forecast_chart(fr)
    if chart_result:
        chart_fig, chart_path = chart_result
        panel["chart"] = gr.update(value=chart_fig, visible=True)
        panel["chart_download"] = gr.update(value=chart_path, visible=True)

    # ---- Traçabilité : modèles comparés ----
    trace = build_trace(result, elapsed)

    ranking = getattr(fr, "candidates", None) or []
    if ranking:
        rows = "\n".join(
            f"| {r['model']} | {format_value(r['MAE'], unit)} | {r['MAPE']:.1f} % |".replace(".", ",")
            for r in ranking
        )
        trace += (
            "\n\n**Modèles comparés (backtest)**\n\n"
            "| Modèle | Erreur moyenne (MAE) | MAPE |\n"
            "|---|---:|---:|\n" + rows
        )

    panel["trace"] = trace

    trace_text = panel.get("trace", "")

    if isinstance(trace_text, str):
        panel["glossary"] = build_dynamic_glossary(trace_text)
    else:
        panel["glossary"] = build_dynamic_glossary("")

    if notes:
        panel["warnings"] = "\n".join(f"- ⚠️ {n}" for n in notes)

    # follow_ups = [
    #     [f"Prévois {metric} pour les 12 prochains mois"],
    #     [f"{metric.capitalize()} par mois"],
    #     [f"{metric.capitalize()} par région"],
    # ]
    # panel["suggestions"] = gr.update(samples=follow_ups, visible=True)

    return message, panel


# ---------------------------------------------------------
# Construction des sorties
# ---------------------------------------------------------

# Ordre des composants du panneau (doit correspondre à build_app)
PANEL_KEYS = (
    "kpi", "table", "chart", "chart_download", "sql", "trace", "warnings",
    "download", "excel", "glossary",
    "highlights", "plain", "pptx", "chip1", "chip2", "chip3", "alerts",
)


def empty_panel() -> dict[str, Any]:
    hidden = gr.update(value=None, visible=False)

    return {
        "kpi": "",
        "table": hidden,
        "chart": hidden,
        "chart_download": hidden,
        "sql": "",
        "trace": "",
        "warnings": "",
        "download": hidden,
        "excel": hidden,
        "glossary": gr.update(
            value=(
                '<div class="dynamic-glossary-empty">'
                "Les termes utilisés dans la réponse apparaîtront ici."
                "</div>"
            ),
            visible=True,
        ),
        "highlights": gr.update(value="", visible=False),
        "plain": "",
        "pptx": hidden,
        "chip1": gr.update(value=" ", visible=False),
        "chip2": gr.update(value=" ", visible=False),
        "chip3": gr.update(value=" ", visible=False),
        "alerts": gr.update(),   # inchangé : seules l'ouverture et la création d'une alerte le modifient
    }

def panel_list(panel: dict[str, Any]) -> list[Any]:
    return [panel[key] for key in PANEL_KEYS]


def extract_requested_year(question: str) -> int | None:
    match = re.search(r"\b((?:19|20)\d{2})\b", question)
    return int(match.group(1)) if match else None


def build_no_data_message(result) -> str:
    """Message naturel lorsqu'aucune donnée n'est trouvée (pipeline classique)."""
    question = getattr(result, "question", "") or ""

    requested_year = extract_requested_year(question)
    max_year = getattr(result, "data_max_year", None)

    if requested_year is not None and max_year is not None and requested_year > max_year:
        return (
            f"Je n’ai pas trouvé d’information concernant votre demande "
            f"pour **{requested_year}**, car les données disponibles "
            f"s’arrêtent en **{max_year}**."
        )

    if requested_year is not None:
        return (
            f"Je n’ai trouvé aucune information correspondant à votre "
            f"demande pour **{requested_year}**."
        )

    return (
        "Je n’ai trouvé aucune information correspondant à votre demande."
        "\n\nVous pouvez essayer avec une autre période ou un autre critère."
    )


def prepare_dataframe_for_display(df):
    """Noms techniques remplacés par des noms naturels pour l'affichage."""
    if df is None:
        return df

    df = df.copy()
    df = df.drop(columns=[c for c in ["Nom_Mois"] if c in df.columns], errors="ignore")

    return df.rename(
        columns={
            "Annee": "Année",
            "Mois": "Mois",
            "Chiffre_Affaires": "Chiffre d’affaires",
        }
    )


#..................................................................
def clean_chat_analysis(text: str) -> str:
    """Nettoie la réponse LLM destinée au chatbot.

    Le tableau détaillé est affiché séparément dans le panneau Résultats.
    """
    if not text:
        return ""

    text = text.strip()

    # Supprime les blocs de tableaux Markdown
    lines = text.splitlines()
    cleaned = []
    in_table = False

    for line in lines:
        stripped = line.strip()

        # Détection d'une ligne de tableau Markdown
        is_table_line = (
            stripped.startswith("|")
            and stripped.endswith("|")
        )

        if is_table_line:
            in_table = True
            continue

        # Ligne de séparation Markdown : |---|---:|
        if in_table and re.fullmatch(r"\|?[\s:\-|]+\|?", stripped):
            continue

        if in_table:
            # Une ligne normale termine le tableau
            in_table = False

        cleaned.append(line)

    text = "\n".join(cleaned)

    # Nettoyage des lignes vides multiples
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()

###############################################################################
# ---------------------------------------------------------
# Bande « À retenir »
# ---------------------------------------------------------


def build_highlights(df: pd.DataFrame, meta: dict, extra: list | None = None) -> str:
    """Plus forte hausse / baisse / anomalies, calculées en code."""
    measures = measure_columns(df)
    if not measures or len(df) < 2:
        return ""

    col = measures[0]
    unit = detect_unit(col)

    def fmt(v: Any) -> str:
        return format_value(v, unit, compact=True)

    def pct(r: float) -> str:
        return f"{r * 100:+.1f} %".replace(".", ",")

    chips: list[tuple[str, str, str]] = list(extra or [])

    cmp_ = meta.get("compare") or {}

    if is_additive(col, df[col].dropna(), unit):
        total_text = fmt(df[col].sum())
        if cmp_.get("pct") is not None and _norm(cmp_.get("measure")) == _norm(col):
            total_text += " · " + delta_html(cmp_["pct"], cmp_.get("label", "")) \
                .replace("<span", "").replace("</span>", "")  # texte simple
        chips.append(("", "Total", total_text))

    # cmp_ = meta.get("compare") or {}
    by_label = cmp_.get("by_label") or {}

    if by_label:
        k, v = max(by_label.items(), key=lambda kv: kv[1])
        if v > 0:
            chips.append(("up", "▲ Plus forte hausse", f"{k} · {pct(v)} {cmp_.get('label', '')}"))
        k, v = min(by_label.items(), key=lambda kv: kv[1])
        if v < 0:
            chips.append(("down", "▼ Plus forte baisse", f"{k} · {pct(v)} {cmp_.get('label', '')}"))
    else:
        points = _ordered_points(df, col)
        if points and len(points) >= 3:
            moves = [
                (f"{a[0]} → {b[0]}", b[1] / a[1] - 1)
                for a, b in zip(points, points[1:])
                if a[1]
            ]
            if moves:
                name, v = max(moves, key=lambda m: m[1])
                if v > 0:
                    chips.append(("up", "▲ Plus forte hausse", f"{name} · {pct(v)}"))
                name, v = min(moves, key=lambda m: m[1])
                if v < 0:
                    chips.append(("down", "▼ Plus forte baisse", f"{name} · {pct(v)}"))
        else:
            data = df.dropna(subset=[col])
            labels = label_columns(df)
            if labels and len(data) >= 2:
                hi, lo = data.loc[data[col].idxmax()], data.loc[data[col].idxmin()]
                chips.append(("up", "▲ Le plus élevé", f"{format_label(hi[labels[0]])} · {fmt(hi[col])}"))
                chips.append(("down", "▼ Le plus faible", f"{format_label(lo[labels[0]])} · {fmt(lo[col])}"))

    for a in (meta.get("anomalies") or [])[:2]:
        gap = a.get("ecart_pct")
        detail = f" de {abs(gap):.0f} %".replace(".", ",") if gap is not None else ""
        chips.append(
            ("warn", "⚠ Valeur anormale", f"{a.get('periode')} · {a.get('sens', 'écart')}{detail} vs attendu")
        )

    if not chips:
        return ""

    body = "".join(
        f'<div class="hl-chip {css}"><span class="hl-title">{html.escape(t)}</span>'
        f'<span class="hl-text">{html.escape(x)}</span></div>'
        for css, t, x in chips[:4]
    )
    return f'<div class="hl-band"><span class="hl-label">À retenir</span>{body}</div>'


# ---------------------------------------------------------
# Export PowerPoint : une diapositive (titre, graphique, 3 phrases)
# ---------------------------------------------------------


def _chart_points(df: pd.DataFrame, spec: dict):
    measures, labels = measure_columns(df), label_columns(df)
    if not measures or not labels:
        return None

    y = next(
        (c for c in (_resolve_column(df, p) for p in str(spec.get("y") or "").split(",")) if c in measures),
        measures[0],
    )
    points = _ordered_points(df, y)
    if points:
        names, values = zip(*points)
        return y, list(names), [float(v) for v in values], "line"

    x = _resolve_column(df, spec.get("x")) or labels[0]
    top = df.dropna(subset=[y]).sort_values(y, ascending=False).head(10)
    return y, [format_label(v) for v in top[x]], [float(v) for v in top[y]], "bar"


def _sentences(text: str, n: int = 3) -> list[str]:
    t = re.sub(r"\[\[SUITES\]\].*", "", text or "", flags=re.S)
    t = re.sub(r"[*_`>#]", "", t)
    t = re.sub(r"^\s*[-•]\s*", "", t, flags=re.M)
    parts = [p.strip() for p in re.split(r"(?<=[.!?])\s+|\n+", t)]
    return [p for p in parts if len(p) > 25][:n]


def export_pptx(df, spec, title, text, data_until=None) -> str | None:
    """Diapositive 16:9 avec graphique natif (modifiable dans PowerPoint)."""
    try:
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.dml.color import RGBColor
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches, Pt
    except ImportError:
        logger.warning("Export PowerPoint indisponible (pip install python-pptx)")
        return None

    try:
        points = _chart_points(df, spec or {})
        if points is None:
            return None
        value_col, names, values, kind = points

        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        indigo = RGBColor(0x4F, 0x46, 0xE5)

        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.18))
        band.fill.solid()
        band.fill.fore_color.rgb = indigo
        band.line.fill.background()

        def add_text(left, top, width, height, lines, size, bold=False, color="1E1B4B"):
            box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
            tf = box.text_frame
            tf.word_wrap = True
            for i, line in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = line
                p.font.size = Pt(size)
                p.font.bold = bold
                p.font.color.rgb = RGBColor.from_string(color)
                p.space_after = Pt(10)

        add_text(0.5, 0.4, 12.3, 0.9, [title or "Analyse"], 28, bold=True)

        cd = CategoryChartData()
        cd.categories = names
        cd.add_series(pretty_column(value_col), values)
        ctype = XL_CHART_TYPE.LINE_MARKERS if kind == "line" else XL_CHART_TYPE.COLUMN_CLUSTERED
        chart = slide.shapes.add_chart(
            ctype, Inches(0.5), Inches(1.5), Inches(8.2), Inches(5.0), cd
        ).chart
        chart.has_legend = False
        chart.has_title = False
        chart.category_axis.tick_labels.font.size = Pt(10)
        chart.value_axis.tick_labels.font.size = Pt(10)

        plot = chart.plots[0]
        plot.has_data_labels = len(names) <= 12
        if plot.has_data_labels:
            plot.data_labels.font.size = Pt(10)
            plot.data_labels.number_format = "#,##0"
            plot.data_labels.number_format_is_linked = False

        series = plot.series[0]
        if kind == "line":
            series.format.line.color.rgb = indigo
            series.smooth = False
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = indigo

        add_text(9.0, 1.5, 3.9, 0.5, ["À retenir"], 16, bold=True, color="4F46E5")
        add_text(9.0, 2.1, 3.9, 4.3, _sentences(text, 3), 14)

        footer = "Assistant Data AI"
        if data_until:
            footer += f" · Données jusqu'à {data_until}"
        footer += f" · Généré le {time.strftime('%d/%m/%Y')}"
        add_text(0.5, 6.9, 12.3, 0.4, [footer], 10, color="64748B")

        path = Path(tempfile.mkdtemp()) / "analyse.pptx"
        prs.save(path)
        return str(path)

    except Exception:
        logger.exception("Export PowerPoint impossible")
        return None


# ---------------------------------------------------------
# Alertes, filtres, pastilles de questions de suite
# ---------------------------------------------------------


def alerts_update(triggered: list[dict]):
    if not triggered:
        return gr.update(value="", visible=False)

    items = ""
    for a in triggered:
        more = f" (+{a['count'] - len(a['sample'])} autres)" if a["count"] > len(a["sample"]) else ""
        detail = " ; ".join(html.escape(s) for s in a["sample"])
        items += (
            f"<li><b>{html.escape(a['name'])}</b> — {html.escape(a['message'])} : {detail}{more}</li>"
        )

    n = len(triggered)
    return gr.update(
        value=(
            f'<div class="alert-banner"><div class="alert-title">🔔 {n} alerte{"s" if n > 1 else ""} active{"s" if n > 1 else ""}'
            f"</div><ul>{items}</ul></div>"
        ),
        visible=True,
    )


def check_alerts_ui():
    try:
        return alerts_update(get_pipeline().get_agent().check_alerts())
    except Exception:
        logger.exception("Vérification des alertes impossible")
        return gr.update(value="", visible=False)


def load_filters():
    try:
        options = get_pipeline().get_agent().filter_options()
    except Exception:
        logger.exception("Chargement des filtres impossible")
        options = {}
    return tuple(
        gr.update(choices=["Toutes", *options.get(key, [])], value="Toutes")
        for key in ("annee", "region")
    )


def _chips_update(suggestions: list[str] | None) -> dict[str, Any]:
    items = [s for s in (suggestions or []) if s][:3]
    return {
        f"chip{i + 1}": gr.update(
            value=items[i] if i < len(items) else " ", visible=i < len(items)
        )
        for i in range(3)
    }


############################################################################""
#.............................................................................
def build_reply(
    result,
    elapsed: float | None = None,
) -> tuple[str, dict[str, Any]]:
    """Transforme le résultat (agent ou pipeline classique) en réponse + panneau."""
    panel = empty_panel()

    sql = getattr(result, "sql", None)
    if sql:
        panel["sql"] = sql

    if result.status == "clarification_needed":
        question = result.clarification_question or "Pouvez-vous préciser votre demande ?"
        return (
            "🤔 **Pour être sûr de bien vous répondre, "
            f"j’ai besoin d’une précision.**\n\n{question}",
            panel,
        )

    if result.status == "validation_error":
        return (
            "Je n’ai pas pu traiter cette demande de manière fiable. "
            "Pouvez-vous reformuler votre question ou préciser "
            "les informations que vous recherchez ?",
            panel,
        )

    if result.status == "error":
        detail = (getattr(result, "errors", None) or [""])[0]
        message = "Je n’ai pas pu traiter cette demande."
        if detail:
            message += f"\n\n{detail}"
        return message, panel

    if result.status != "ok":
        logger.warning("État inattendu retourné par le pipeline : %s", result.status)
        return (
            "Je rencontre un petit problème pour traiter cette demande. "
            "Vous pouvez réessayer dans un instant.",
            panel,
        )

    if result.warnings:
        panel["warnings"] = "\n".join(f"- ⚠️ {w}" for w in result.warnings)

    analysis_text = getattr(result, "analysis_text", None)
    meta = getattr(result, "metadata", None) or {}

    # ---- Éléments communs : pastilles, phrase en langage clair, alertes ----
    panel.update(_chips_update(meta.get("suggestions")))

    plain = meta.get("plain") or []
    if plain:
        panel["plain"] = "💬 " + " ".join(plain)

    if meta.get("alert_created"):
        try:
            panel["alerts"] = alerts_update(get_pipeline().get_agent().check_alerts())
        except Exception:
            logger.exception("Rafraîchissement des alertes impossible")

    # ---- Prévision ----
    if getattr(result, "forecast_result", None) is not None:
        message, panel = build_forecast_reply(result, panel, elapsed)
        final_message = clean_chat_analysis(analysis_text) if analysis_text else message
        panel["glossary"] = gr.update(
            value=build_dynamic_glossary(f"{panel['trace']}\n{final_message}\n{sql or ''}"),
            visible=True,
        )
        return final_message, panel

    # ---- Données ----
    dataframe = result.dataframe
    if dataframe is not None and not dataframe.empty:
        dataframe = dataframe.dropna(how="all")

    if dataframe is None or dataframe.empty:
        if analysis_text:
            return clean_chat_analysis(analysis_text), panel
        return build_no_data_message(result), panel

    df = clean_dataframe(dataframe)
    rows = len(df)
    question = getattr(result, "question", "") or ""

    if analysis_text:
        message = clean_chat_analysis(analysis_text)
    else:
        answer = build_answer(question, df)
        message = answer or "Voici ce que j’ai trouvé."
        insights = build_insights(question, df) if rows > 1 else []
        if insights:
            message += "\n\n**📌 L'essentiel**\n" + "\n".join(f"- {line}" for line in insights)
        if rows > 1 or not answer:
            message += f"\n\n> {plural(rows, 'résultat')} — détail dans le tableau."

    panel["kpi"] = build_kpi_html(
        df,
        elapsed,
        question,
        last_updated=getattr(result, "last_updated", None),
        compare=meta.get("compare"),
        data_until=meta.get("data_until"),
    )

    display_df = prepare_dataframe_for_display(df)
    panel["table"] = gr.update(value=style_dataframe(display_df), visible=True)
    panel["download"] = gr.update(value=export_csv(display_df), visible=True)

    excel_path = export_excel(display_df)
    panel["excel"] = gr.update(value=excel_path, visible=bool(excel_path))

    chart_spec = dict(meta.get("chart") or {})
    if not chart_spec.get("years"):
        chart_spec["years"] = _years_from(getattr(result, "sql", None), question)

    focus = meta.get("focus_region")
    regions_df = meta.get("regions_df")

    chart, extra = None, []
    if focus and regions_df is not None:
        rdf = clean_dataframe(regions_df)
        chart = _focus_map(rdf, focus, chart_spec["years"])
        chip = _rank_chip(rdf, focus)
        if chip:
            extra.append(chip)

    if chart is None:
        chart = make_chart_from_spec(df, chart_spec, meta.get("anomalies"))

    panel["chart"] = gr.update(value=chart, visible=chart is not None)
    panel["chart_download"] = gr.update(value=None, visible=False)

    highlights = build_highlights(df, meta, extra)
    panel["highlights"] = gr.update(value=highlights, visible=bool(highlights))

    slide_title = chart.layout.title.text if chart is not None else question
    pptx_path = export_pptx(df, chart_spec, slide_title, message, meta.get("data_until"))
    panel["pptx"] = gr.update(value=pptx_path, visible=bool(pptx_path))

    panel["trace"] = build_trace(result, elapsed)
    panel["glossary"] = gr.update(
        value=build_dynamic_glossary(f"{panel['trace']}\n{message}\n{sql or ''}"),
        visible=True,
    )

    return message, panel


# ---------------------------------------------------------
# Callbacks
# ---------------------------------------------------------


def handle_question(
    message: str,
    chat: list[dict],
    questions: list[str],
    memory: list[dict],
    year: str = "Toutes",
    region: str = "Toutes",
    focus_region: str | None = None,
):
    """Traite une question. Générateur : affiche d'abord un état d'attente."""
    message = (message or "").strip()

    def outputs(chat_value, panel, questions_value, memory_value):
        return (
            chat_value,
            "",
            gr.update(interactive=False),
            *panel_list(panel),
            questions_value,
            gr.update(samples=[[q] for q in reversed(questions_value[-8:])]),
            memory_value,
        )

    if not message:
        gr.Warning("Saisissez une question avant d'envoyer.")
        yield outputs(chat, empty_panel(), questions, memory)
        return

    filters = {
        label: value
        for label, value in (("Année", year), ("Région", region))
        if value and value != "Toutes"
    }
    if focus_region:
        filters.pop("Région", None)

    history = [
        m for m in chat
        if isinstance(m.get("content"), str)
        and m["content"] != WELCOME
        and "chat-loading" not in m["content"]
    ]

    chat = chat + [
        {"role": "user", "content": message},
        {
            "role": "assistant",
            "content": """
            <div class="chat-loading">
                <span class="loading-hourglass">⏳</span>
                <span>Analyse des données en cours…</span>
            </div>
            """,
        },
    ]

    yield outputs(chat, empty_panel(), questions, memory)

    try:
        start = time.perf_counter()
        result = get_pipeline().run_chat(message, history, memory, filters, focus_region)
        elapsed = time.perf_counter() - start
        reply, panel = build_reply(result, elapsed)

        memo = (getattr(result, "metadata", None) or {}).get("memo")
        if memo and memo.get("sql"):
            memory = (memory + [memo])[-5:]

    except Exception:
        logger.exception("Erreur pendant l'exécution du pipeline")
        reply = (
            "Je rencontre actuellement un problème pour récupérer "
            "ces informations. Vous pouvez réessayer dans un instant."
        )
        panel = empty_panel()

    chat[-1] = {"role": "assistant", "content": reply}

    if message not in questions:
        questions = questions + [message]

    yield outputs(chat, panel, questions, memory)


def reset_all():
    return (
        [{"role": "assistant", "content": WELCOME}],
        "",
        *panel_list(empty_panel()),
        [],   # mémoire des requêtes
    )


####################################""

N_OUTPUTS = 6 + len(PANEL_KEYS)   # chatbot, msg, bouton, panneau, questions, historique, mémoire


def _match_region(name: Any) -> str | None:
    """Nom reconnu parmi les régions proposées dans le filtre (ou None)."""
    if not name:
        return None
    try:
        options = get_pipeline().get_agent().filter_options().get("region", [])
    except Exception:
        return None
    key = _plain_key(name)
    return next((o for o in options if _plain_key(o) == key), None)


ZOOM_PREFIX = "Détail du chiffre d'affaires de la région"


def _user_messages(chat: list[dict]) -> list[str]:
    return [
        m["content"] for m in chat
        if m.get("role") == "user" and isinstance(m.get("content"), str)
    ]


def refilter(chat, memory, questions, year, region):
    """Année ou Région change : on relance la dernière question avec les filtres."""
    users = _user_messages(chat)
    last = users[-1] if users else None
    last_real = next((q for q in reversed(users) if not q.startswith(ZOOM_PREFIX)), None)

    zoom_active = last is not None and last.startswith(ZOOM_PREFIX)

    if zoom_active and region != "Toutes":
        yield from handle_question(
            f"{ZOOM_PREFIX} {region}", chat, questions, memory, year, region, focus_region=region
        )
    elif last_real:
        yield from handle_question(last_real, chat, questions, memory, year, region)
    elif region != "Toutes":
        yield from handle_question(
            f"{ZOOM_PREFIX} {region}", chat, questions, memory, year, region, focus_region=region
        )
    else:
        yield tuple(gr.update() for _ in range(N_OUTPUTS))


def select_region(name, chat, memory, questions, year):
    """Clic (carte ou tableau) : met à jour la liste Région puis lance le zoom."""
    match = _match_region(name)
    if not match:
        yield tuple(gr.update() for _ in range(N_OUTPUTS + 1))
        return

    for out in handle_question(
        f"{ZOOM_PREFIX} {match}", chat, questions, memory, year, match, focus_region=match
    ):
        yield (gr.update(value=match), *out)


def on_map_click(name, chat, memory, questions, year):
    yield from select_region(name, chat, memory, questions, year)


def on_table_select(evt: gr.SelectData, chat, memory, questions, year):
    cells = getattr(evt, "row_value", None) or [evt.value]
    name = next((c for c in cells if _match_region(c)), None)
    yield from select_region(name, chat, memory, questions, year)

#....................................................

GLOSSARY = {
    "MAPE": (
        "Erreur moyenne en pourcentage entre les valeurs réelles "
        "et les prévisions. Plus elle est faible, mieux c'est."
    ),
    "MAE": (
        "Erreur absolue moyenne entre les valeurs réelles "
        "et les prévisions. Plus elle est faible, mieux c'est."
    ),
    "RMSE": (
        "Erreur qui pénalise davantage les grosses erreurs. "
        "Plus elle est faible, mieux c'est."
    ),
    "XGBoost": (
        "Algorithme de machine learning basé sur des arbres de décision, "
        "capable de détecter des relations complexes dans les données."
    ),
    "Holt-Winters": (
        "Méthode statistique de prévision qui prend en compte "
        "la tendance et la saisonnalité."
    ),
    "Holt amorti": (
        "Méthode de prévision qui modélise la tendance "
        "en réduisant progressivement son influence dans le futur."
    ),
    "Régression linéaire": (
        "Méthode statistique qui estime une tendance entre "
        "une variable à prévoir et des variables explicatives."
    ),
    "Régression linéaire saisonnière": (
        "Régression linéaire qui tient également compte "
        "des variations saisonnières."
    ),
    "Saisonnalité": (
        "Variation qui se répète régulièrement dans le temps, "
        "par exemple chaque mois ou chaque année."
    ),
    "Backtesting": (
        "Technique consistant à tester un modèle sur des périodes "
        "historiques connues afin d'évaluer ses performances."
    ),
    "Forecast": (
        "Prévision des valeurs futures à partir des données historiques."
    ),
    "Prévision": (
        "Estimation de valeurs futures à partir de l'historique disponible."
    ),
    "Modèle sélectionné": (
        "Modèle ayant obtenu les meilleures performances "
        "lors des tests historiques."
    ),
}


######################################################"
BUSINESS_GLOSSARY = {
    "Chiffre d'affaires": "Somme des montants des ventes sur la période (colonne Chiffre_Affaires de Fact_Ventes).",
    "CA": "Abréviation de chiffre d'affaires.",
    "Marge": "Somme des marges des ventes sur la période (colonne Marge de Fact_Ventes).",
    "Taux de marge": "Marge divisée par le chiffre d'affaires de la même période, exprimée en %.",
    "Quantité": "Nombre d'unités vendues (colonne Quantite de Fact_Ventes).",
    "Panier moyen": "Chiffre d'affaires divisé par le nombre de ventes.",
    "N-1": "Même période de l'année précédente, utilisée comme base de comparaison.",
    "Anomalie": (
        "Valeur statistiquement inhabituelle compte tenu de la tendance "
        "et de la saisonnalité (écart supérieur à 3 écarts robustes)."
    ),
    "Tendance": "Évolution moyenne par mois, calculée après neutralisation de la saisonnalité.",
}
GLOSSARY.update(BUSINESS_GLOSSARY)


def build_dynamic_glossary(text: str) -> str:
    def empty(msg: str) -> str:
        return f'<div class="dynamic-glossary-empty">{msg}</div>'

    if not text:
        return empty("Les termes utilisés dans la réponse apparaîtront ici.")

    text = text.replace("’", "'")
    found = []

    for term, definition in GLOSSARY.items():
        flags = 0 if len(term) <= 3 else re.IGNORECASE   # « CA », « MAE » : casse stricte
        m = re.search(rf"\b{re.escape(term)}\b", text, flags)
        if m:
            found.append((m.start(), m.end(), term, definition))

    # « Marge » est masqué quand « Taux de marge » le contient
    found = [
        f for f in found
        if not any(
            o is not f and o[0] <= f[0] and o[1] >= f[1] and (o[1] - o[0]) > (f[1] - f[0])
            for o in found
        )
    ]
    found.sort(key=lambda x: x[0])

    if not found:
        return empty("Aucun terme technique ou métier détecté.")

    blocks = "".join(
        f'<div class="glossary-item"><div class="glossary-term">{html.escape(term)}</div>'
        f'<div class="glossary-definition">{html.escape(definition)}</div></div>'
        for _, _, term, definition in found
    )
    return f'<div class="dynamic-glossary">{blocks}</div>'




#.......................................


# ---------------------------------------------------------
# Interface
# ---------------------------------------------------------


def _sidebar():
    """Menu latéral à droite (repli sur la gauche si Gradio est ancien)."""
    try:
        return gr.Sidebar(position="right", open=True, width=300)
    except TypeError:
        return gr.Sidebar(open=True, width=300)


############################"


def build_about_html() -> str:
    return """
    <style>
    #about-page{max-width:1250px;margin:auto;padding:10px 10px 40px;color:#1e1b4b}
    .about-hero,.about-card,.about-section{background:#fff;border:1px solid #e0e7ff;border-radius:16px;box-shadow:0 2px 8px rgba(79,70,229,.06)}
    .about-hero{display:flex;align-items:center;gap:18px;padding:24px 28px;margin-bottom:14px;background:linear-gradient(135deg,#eef2ff,#e0f2fe)}
    .about-hero-icon{width:58px;height:58px;display:flex;align-items:center;justify-content:center;background:#fff;border:1px solid #c7d2fe;border-radius:15px;font-size:28px}
    .about-hero h1{margin:0 0 5px;font-size:1.7rem}.about-hero p{margin:0;color:#5b5f72;font-size:.95rem}
    .about-grid{display:grid;grid-template-columns:repeat(3,1fr);gap:14px;margin-bottom:14px}
    .about-card{padding:20px}.about-card-icon{font-size:25px;margin-bottom:8px}
    .about-card h2,.about-section h2{margin:0 0 8px;font-size:1.05rem}
    .about-card p{margin:7px 0;line-height:1.55;font-size:.88rem;color:#4b5563}
    .about-card ul,.about-limits ul{margin:10px 0 0;padding-left:19px;color:#4b5563;font-size:.86rem;line-height:1.65}
    .about-flow-mini{display:flex;flex-wrap:wrap;gap:6px;margin-top:15px;padding:10px;background:#f8fafc;border-radius:10px;color:#3730a3;font-size:.78rem;font-weight:600}
    .about-section{padding:22px;margin-bottom:14px}
    .about-section-title{display:flex;align-items:center;gap:12px;margin-bottom:18px}.about-section-title>span{font-size:25px}
    .about-section-title p{margin:0;color:#6b7280;font-size:.82rem}
    .about-process{display:flex;align-items:stretch;gap:8px}.about-step{flex:1;padding:14px;background:#f8fafc;border:1px solid #e5e7eb;border-radius:12px}
    .about-step-number{width:27px;height:27px;display:flex;align-items:center;justify-content:center;margin-bottom:9px;background:#4f46e5;color:#fff;border-radius:50%;font-size:.75rem;font-weight:700}
    .about-step strong{display:block;font-size:.85rem;margin-bottom:5px}.about-step span{color:#6b7280;font-size:.75rem;line-height:1.4}.about-arrow{display:flex;align-items:center;color:#a5b4fc;font-size:20px}
    .about-two-columns{display:grid;grid-template-columns:1fr 1fr;gap:14px}.about-value-list{display:grid;gap:9px}
    .about-value-list>div{display:flex;flex-direction:column;padding:11px 13px;background:#f8fafc;border-radius:10px}
    .about-value-list b{font-size:.82rem;margin-bottom:3px}.about-value-list span{color:#6b7280;font-size:.77rem;line-height:1.4}
    .about-warning{padding:12px 15px;margin-bottom:12px;background:#fff7ed;border:1px solid #fdba74;border-left:4px solid #f97316;border-radius:10px;color:#9a3412;font-size:.83rem}.about-warning p{margin:4px 0}
    .about-transparency-flow{display:flex;align-items:center;gap:9px;padding:15px;background:#f8fafc;border-radius:12px;overflow-x:auto}
    .about-transparency-flow>div{min-width:125px;padding:11px;background:#fff;border:1px solid #e0e7ff;border-radius:10px;text-align:center}
    .about-transparency-flow span{display:block;font-weight:700;font-size:.78rem;color:#3730a3}.about-transparency-flow small{display:block;margin-top:4px;color:#6b7280;font-size:.68rem}.about-transparency-flow b{color:#a5b4fc}
    .about-footer-note{margin:15px 0 0;padding-top:13px;border-top:1px solid #e5e7eb;color:#6b7280;font-size:.8rem;line-height:1.5}
    @media(max-width:900px){.about-grid,.about-two-columns{grid-template-columns:1fr}.about-process{flex-direction:column}.about-arrow{justify-content:center;transform:rotate(90deg)}}
    </style>
    <div id="about-page">
      <div class="about-hero"><div class="about-hero-icon">ℹ️</div><div><h1>À propos de l'Assistant</h1><p>Comprendre le rôle, le fonctionnement et la valeur d'Assistant Data AI dans l'entreprise.</p></div></div>
      <div class="about-grid">
        <div class="about-card"><div class="about-card-icon">🎯</div><h2>Objectif de l'Assistant</h2><p>Assistant Data AI permet d'interroger les données de l'entreprise en langage naturel, sans avoir besoin de maîtriser SQL ou les outils de reporting.</p><p>L'objectif est de rendre l'information plus accessible, plus rapide à obtenir et plus simple à exploiter.</p></div>
        <div class="about-card"><div class="about-card-icon">💼</div><h2>Le côté métier</h2><p>L'Assistant transforme une question métier en une analyse exploitable à travers des indicateurs, des tableaux et des visualisations.</p><ul><li>Suivre les indicateurs d'activité</li><li>Comparer des périodes ou des périmètres</li><li>Analyser les régions et les produits</li><li>Identifier des évolutions et des écarts</li><li>Préparer rapidement une analyse métier</li></ul></div>
        <div class="about-card"><div class="about-card-icon">📊</div><h2>Les données au cœur de l'analyse</h2><p>L'Assistant s'appuie sur les données disponibles dans l'environnement de l'entreprise afin de produire des résultats cohérents avec le périmètre accessible.</p><div class="about-flow-mini"><span>Question</span><span>→</span><span>Données</span><span>→</span><span>Analyse</span><span>→</span><span>Résultat</span></div></div>
      </div>
      <div class="about-section"><div class="about-section-title"><span>🔄</span><div><h2>Comment ça fonctionne ?</h2><p>De la question métier jusqu'au résultat.</p></div></div><div class="about-process">
        <div class="about-step"><div class="about-step-number">1</div><strong>Question métier</strong><span>L'utilisateur formule sa demande en langage naturel.</span></div><div class="about-arrow">→</div>
        <div class="about-step"><div class="about-step-number">2</div><strong>Compréhension</strong><span>L'Assistant interprète la demande et son contexte.</span></div><div class="about-arrow">→</div>
        <div class="about-step"><div class="about-step-number">3</div><strong>Requête</strong><span>La demande est traduite en interrogation des données.</span></div><div class="about-arrow">→</div>
        <div class="about-step"><div class="about-step-number">4</div><strong>Analyse</strong><span>Les résultats sont analysés et mis en forme.</span></div><div class="about-arrow">→</div>
        <div class="about-step"><div class="about-step-number">5</div><strong>Restitution</strong><span>KPI, tableau, graphique et réponse en langage clair.</span></div>
      </div></div>
      <div class="about-two-columns">
        <div class="about-section"><div class="about-section-title"><span>🚀</span><div><h2>Valeur ajoutée pour l'entreprise</h2><p>Pourquoi utiliser Assistant Data AI ?</p></div></div><div class="about-value-list">
          <div><b>⏱ Gain de temps</b><span>Accéder rapidement à une information sans construire manuellement une requête.</span></div>
          <div><b>🎯 Autonomie métier</b><span>Permettre aux équipes métier d'explorer les données plus facilement.</span></div>
          <div><b>📊 Accessibilité de la donnée</b><span>Rendre les données accessibles aux utilisateurs non techniques.</span></div>
          <div><b>🔎 Aide à la décision</b><span>Faire ressortir rapidement tendances, écarts et indicateurs importants.</span></div>
          <div><b>💡 Culture Data</b><span>Encourager une utilisation plus régulière des données dans les décisions.</span></div>
        </div></div>
        <div class="about-section about-limits"><div class="about-section-title"><span>⚠️</span><div><h2>Limites et bonnes pratiques</h2><p>Un assistant d'analyse, pas une vérité absolue.</p></div></div>
          <div class="about-warning"><strong>Assistant Data AI est un outil d'aide à l'analyse.</strong><p>Il ne remplace pas l'expertise métier ni la validation des données.</p></div>
          <ul><li>Les réponses dépendent de la qualité des données.</li><li>Une question ambiguë peut être interprétée différemment.</li><li>Les résultats doivent être replacés dans leur contexte métier.</li><li>Une décision importante doit être vérifiée.</li><li>L'Assistant ne travaille que sur les données auxquelles il a accès.</li><li>Une absence de résultat ne signifie pas nécessairement une absence d'activité.</li></ul>
        </div>
      </div>
      <div class="about-section"><div class="about-section-title"><span>🔎</span><div><h2>Transparence et confiance</h2><p>Comprendre comment une réponse est construite.</p></div></div>
        <div class="about-transparency-flow"><div><span>Question</span><small>Demande utilisateur</small></div><b>→</b><div><span>Filtres</span><small>Année · Région · Produit</small></div><b>→</b><div><span>Requête</span><small>Interrogation des données</small></div><b>→</b><div><span>Résultats</span><small>Données retournées</small></div><b>→</b><div><span>Visualisation</span><small>KPI · Tableau · Graphique</small></div></div>
        <p class="about-footer-note">L'objectif est de proposer une analyse compréhensible, contextualisée et exploitable par les utilisateurs métier.</p>
      </div>
    </div>
    """

def build_app() -> gr.Blocks:
    with gr.Blocks(title="Assistant Data AI") as app:
        questions_state = gr.State([])
        memory_state = gr.State([])

        # ---------------- En-tête ----------------
        with gr.Row(elem_id="entete", equal_height=True):
            with gr.Row(elem_id="entete-gauche", scale=1, equal_height=True):
                if HAS_LOGO:
                    gr.Image(
                        value=str(LOGO_PATH),
                        height=60,
                        width=60,
                        show_label=False,
                        container=False,
                        interactive=False,
                        buttons=[],
                        scale=0,
                        min_width=60,
                        elem_id="logo",
                    )

                gr.Markdown(
                    "# Assistant Data AI\nInterrogez vos données en langage naturel.",
                    elem_id="titre",
                )
            kpi = gr.HTML(elem_id="kpi", scale=5)

            with gr.Row(elem_id="header-actions"):
                about_btn = gr.Button(
                    "ℹ️ À propos de l'Assistant Data AI",
                    elem_id="about-button",
                    variant="secondary",
                    size="sm",
                    scale=0,
                    min_width=0,
                )

        # ---------------- Page « À propos » ----------------
# ---------------- Page « À propos » ----------------
        with gr.Column(
            visible=False,
            elem_id="about-page-container"
        ) as about_page:

            back_btn = gr.Button(
                "← Retour à l'Assistant",
                elem_id="about-back-button",
                variant="secondary",
                size="sm",
                scale=0,
                min_width=0,
            )

            gr.HTML(
                value=build_about_html(),
                elem_id="about-content",
            )
            #........................

        # about_close_trigger = gr.Button(
        #     "Retour",
        #     elem_id="about-close-trigger",
        #     visible=False,
        # )

        # about_close_trigger.click(
        #     fn=lambda: gr.update(visible=False),
        #     inputs=None,
        #     outputs=about_page,
        # )

        # ---------------- Alertes (visibles seulement si déclenchées) ----------------
        alerts_html = gr.HTML(visible=False, elem_id="alerts")

        # # ---------------- Filtres globaux ----------------
        # with gr.Row(elem_id="filtres", equal_height=True):
        #     year_dd = gr.Dropdown(["Toutes"], value="Toutes", label="Année", scale=1, min_width=140)
        #     region_dd = gr.Dropdown(["Toutes"], value="Toutes", label="Région", scale=2, min_width=200)
        #     product_dd = gr.Dropdown(["Toutes"], value="Toutes", label="Produit", scale=2, min_width=200)
        #     map_click = gr.Textbox(elem_id="map-click", container=False, label="map")

        # ---------------- Contenu principal ----------------
        with gr.Row(elem_id="workspace", equal_height=False):

            with gr.Column(scale=4, min_width=320, elem_id="chat-column"):
                chatbot = gr.Chatbot(
                    value=[{"role": "assistant", "content": WELCOME}],
                    height="36vh",
                    show_label=False,
                    render_markdown=True,
                    avatar_images=(None, str(LOGO_PATH) if HAS_LOGO else None),
                    elem_id="chatbot",
                )

                with gr.Row(elem_id="saisie"):
                    msg = gr.Textbox(
                        placeholder="Ex. : Quel est le chiffre d'affaires par région ?",
                        show_label=False, container=False, lines=2, max_lines=4,
                        autofocus=True, scale=6,
                    )
                    send_btn = gr.Button(
                        "Analyser", variant="primary", scale=2, min_width=110, interactive=False
                    )
                    clear_btn = gr.Button("↺ Nouveau", scale=1, min_width=90)

                    msg.input(
                        lambda text: gr.update(interactive=bool(text and text.strip())),
                        inputs=msg,
                        outputs=send_btn,
                    )

                with gr.Row(elem_id="chips"):
                    chip1 = gr.Button(" ", visible=False, size="sm", elem_classes="chip")
                    chip2 = gr.Button(" ", visible=False, size="sm", elem_classes="chip")
                    chip3 = gr.Button(" ", visible=False, size="sm", elem_classes="chip")

                gr.Markdown("### Requête SQL", elem_id="titre-sql")
                sql_box = gr.Code(
                    language="sql", interactive=False, show_label=False,
                    lines=6, max_lines=10, elem_id="sql-box",
                )
                # Phrase en langage clair, APRÈS la requête SQL
                plain_md = gr.Markdown(elem_id="plain-sql")

                # gr.Markdown("### Requête SQL", elem_id="titre-sql")
                # sql_box = gr.Code(
                #     language="sql", interactive=False, show_label=False, elem_id="sql-box"
                # )

            with gr.Column(scale=8, min_width=600, elem_id="results-column"):
                with gr.Group(elem_id="graphique-principal"):
                    gr.Markdown("### Visualisation", elem_id="titre-graphique")

                    # Bande « À retenir » au-dessus du graphique
                    highlights_html = gr.HTML(visible=False, elem_id="highlights")

                    with gr.Row(elem_id="chart-row", equal_height=True):
                        chart = gr.Plot(
                            visible=False, show_label=False, elem_id="chart",
                            scale=12, min_width=400,
                        )
                        chart_download = gr.DownloadButton(
                            "⬇", variant="secondary", visible=False, scale=0,
                            min_width=42, elem_id="chart-download",
                        )

                table = gr.Dataframe(
                    visible=False, interactive=False, wrap=True,
                    max_height="30vh", elem_id="table", show_label=False,
                )

                with gr.Row(elem_id="download-row"):
                    download = gr.DownloadButton(
                        "Télécharger en CSV", variant="primary", visible=False
                    )
                    excel = gr.DownloadButton(
                        "Télécharger en Excel", variant="secondary", visible=False
                    )
                    pptx_btn = gr.DownloadButton(
                        "Télécharger la diapositive (PowerPoint)", variant="secondary", visible=False
                    )

        # ---------------- Informations techniques ----------------
        with gr.Tabs(elem_id="technical-tabs"):
            with gr.Tab("Avertissements"):
                warnings_md = gr.Markdown(visible=True)
            with gr.Tab("Traçabilité"):
                trace_md = gr.Markdown(visible=True)

        # ---------------- Menu latéral ----------------
        with _sidebar():
            with gr.Column(elem_id="assistant-sidebar-content"):
                gr.Markdown("### 🎛️ Filtres", elem_id="filtres-title")
                with gr.Column(elem_id="filtres"):
                    year_dd = gr.Dropdown(["Toutes"], value="Toutes", label="Année")
                    region_dd = gr.Dropdown(["Toutes"], value="Toutes", label="Région")
                map_click = gr.Textbox(elem_id="map-click", container=False, label="map")

                gr.Markdown("### 📚 Mots-clés", elem_id="glossary-title")
                glossary_html = gr.HTML(
                    value=(
                        '<div class="dynamic-glossary-empty">'
                        "Les termes utilisés dans la réponse apparaîtront ici."
                        "</div>"
                    ),
                    elem_id="dynamic-glossary-container",
                )
                history_ds = gr.Dataset(
                    components=[gr.Textbox(visible=False)],
                    samples=[],
                    label="Historique récent",
                    type="values",
                )

        # L'ordre DOIT rester identique à PANEL_KEYS
        panel_components = [
            kpi, table, chart, chart_download, sql_box, trace_md, warnings_md,
            download, excel, glossary_html,
            highlights_html, plain_md, pptx_btn, chip1, chip2, chip3, alerts_html,
        ]

        # ---------------- Événements ----------------
        ABOUT_OPEN_JS = """() => {
            document.body.classList.add('about-mode');
        }"""

        ABOUT_CLOSE_JS = """() => {
            document.body.classList.remove('about-mode');
        }"""

        about_btn.click(
            fn=lambda: gr.update(visible=True),
            inputs=None,
            outputs=about_page,
            js=ABOUT_OPEN_JS,
        )

        back_btn.click(
            fn=lambda: gr.update(visible=False),
            inputs=None,
            outputs=about_page,
            js=ABOUT_CLOSE_JS,
        )

        back_btn.click(
            fn=lambda: gr.update(visible=False),
            inputs=None,
            outputs=about_page,
            js=ABOUT_CLOSE_JS,
        )
        # def show_assistant():
        #     return (
        #         gr.update(visible=True),
        #         gr.update(visible=False),
        #     )

        # back_btn.click(
        #     fn=show_assistant,
        #     inputs=None,
        #     outputs=[
        #         main_content,
        #         about_page,
        #     ],
        # )

        submit_inputs = [msg, chatbot, questions_state, memory_state, year_dd, region_dd]
        submit_outputs = [
            chatbot, msg, send_btn,
            *panel_components,
            questions_state, history_ds, memory_state,
        ]

        send_btn.click(handle_question, submit_inputs, submit_outputs)
        msg.submit(handle_question, submit_inputs, submit_outputs)
        
########################""
        filter_inputs = [chatbot, memory_state, questions_state, year_dd, region_dd]

        for dropdown in (year_dd, region_dd):
            dropdown.input(refilter, filter_inputs, submit_outputs)

        click_inputs = [chatbot, memory_state, questions_state, year_dd]
        click_outputs = [region_dd, *submit_outputs]

        map_click.input(on_map_click, [map_click, *click_inputs], click_outputs)
        table.select(on_table_select, click_inputs, click_outputs)

        app.load(None, None, None, js=MAP_CLICK_JS)


        for chip in (chip1, chip2, chip3):
            chip.click(lambda v: v, chip, msg).then(handle_question, submit_inputs, submit_outputs)

        clear_btn.click(
            reset_all, None, [chatbot, msg, *panel_components, memory_state],
        ).then(lambda: gr.update(interactive=False), None, send_btn)

        history_ds.click(lambda row: row[0], history_ds, msg)

        # À l'ouverture : options des filtres + vérification des alertes
        app.load(load_filters, None, [year_dd, region_dd])
        app.load(check_alerts_ui, None, alerts_html)

        gr.HTML("""
        <style>
            body.about-mode #workspace,
            body.about-mode #technical-tabs,
            body.about-mode #alerts,
            body.about-mode #assistant-sidebar-content {
                display: none !important;
            }
            body.about-mode #about-page-container {
                display: block !important;
            }
            body.about-mode #about-button {
                display: none !important;
            }
            #about-page-container { width: 100%; }
            #header-actions { justify-content: flex-end; align-items: center; }
            #about-back-button { margin: 0 0 12px 0; }
        </style>
        """)

    return app



if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)

    # Préchauffage : connexion SQL, client OpenAI et catalogue chargés au démarrage
    try:
        get_pipeline()
    except Exception:
        logger.exception("Préchauffage du pipeline impossible")

    build_app().queue().launch(
        favicon_path=str(LOGO_PATH) if HAS_LOGO else None,
        allowed_paths=[str(IMAGE_DIR)],
        theme=THEME,
        css=CSS,
        js=FORCE_LIGHT_JS,
    )